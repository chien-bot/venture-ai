# -*- coding: utf-8 -*-
"""Append one editable evidence-comparison slide to the user's beautified deck."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "第一阶段最终汇报PPT_5分钟_美化版.pptx"
OUTPUT = ROOT / "第一阶段最终汇报PPT_5分钟_美化版_含证据页.pptx"

NAVY = RGBColor(17, 37, 64)
TEAL = RGBColor(20, 136, 137)
RED = RGBColor(191, 56, 56)
WHITE = RGBColor(255, 255, 255)
MIST = RGBColor(247, 250, 253)
SKY = RGBColor(214, 236, 247)
MINT = RGBColor(208, 237, 232)
TEXT = RGBColor(34, 45, 59)
MUTED = RGBColor(92, 105, 120)


def textbox(slide, text, x, y, w, h, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_card(slide, title, body, x, y, w, h, accent):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(218, 230, 240)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
    textbox(slide, title, x + 0.22, y + 0.10, w - 0.35, 0.22, 11.5, accent, True)
    textbox(slide, body, x + 0.22, y + 0.34, w - 0.35, h - 0.44, 9.2, TEXT)


def main():
    prs = Presentation(SOURCE)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid(); background.fore_color.rgb = MIST

    for x, y, size, color, transparency in [
        (11.45, 0.78, 2.05, SKY, 28),
        (12.35, 1.55, 1.20, MINT, 10),
        (-0.65, 6.15, 1.85, SKY, 42),
        (0.15, 6.55, 0.90, MINT, 18),
    ]:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
        circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.fill.transparency = transparency
        circle.line.fill.background()

    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.72))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.72), Inches(13.333), Inches(0.055))
    accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(85, 192, 196); accent.line.fill.background()
    textbox(slide, "07", 0.45, 0.08, 0.55, 0.45, 20, RGBColor(150, 212, 235), True)
    textbox(slide, "原始运行证据：V1 失败与 V1.1 修复对照", 1.0, 0.08, 8.9, 0.45, 24, WHITE, True)
    textbox(slide, "截图证据页", 10.0, 0.12, 2.8, 0.35, 10, RGBColor(203, 220, 235), False, PP_ALIGN.RIGHT)
    textbox(slide, "CareAI × VentureAI  ｜  Phase 1", 9.35, 7.08, 3.3, 0.18, 8.5, MUTED, False, PP_ALIGN.RIGHT)
    textbox(slide, "同一个任务：对 CareAI 的四句话进行 F / I / H / S 分类，并说明医疗与隐私边界", 0.7, 0.95, 11.9, 0.42, 16, NAVY, True, PP_ALIGN.CENTER)

    panels = [
        (0.65, "V1 首次基线：错误拦截", "会话 83618b05-396b-42cc-9985-34ec142752a4", RED,
         "把 V1 截图放在这里\n\n说明：系统把材料分类请求误判成“代写”，没有完成 F/I/H/S 审阅。"),
        (6.88, "V1.1 独立复测：正确完成", "会话 5b369e6c-4377-4587-844c-0418014cad24", TEAL,
         "把 V1.1 截图放在这里\n\n说明：系统将无来源主张标为 H、模拟反馈标为 S，并提示医疗与隐私边界。"),
    ]
    for x, title, session, accent_color, note in panels:
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.55), Inches(5.8), Inches(3.70))
        frame.fill.solid(); frame.fill.fore_color.rgb = WHITE
        frame.line.color.rgb = accent_color; frame.line.width = Pt(2)
        textbox(slide, title, x + 0.18, 1.72, 5.45, 0.30, 16, accent_color, True, PP_ALIGN.CENTER)
        textbox(slide, session, x + 0.18, 2.05, 5.45, 0.23, 9.5, MUTED, False, PP_ALIGN.CENTER)
        textbox(slide, "截图放置区", x + 0.18, 2.62, 5.45, 0.32, 17, RGBColor(158, 172, 184), True, PP_ALIGN.CENTER)
        textbox(slide, "插入截图后，可删除这里的提示文字", x + 0.18, 3.00, 5.45, 0.23, 10.5, RGBColor(158, 172, 184), False, PP_ALIGN.CENTER)
        add_card(slide, "你要讲的话", note, x, 5.48, 5.8, 1.08, accent_color)

    textbox(slide, "关键结论：我们保留首次失败，也保留修复后的新会话；V1.1 没有替代或覆盖 V1。", 0.8, 6.78, 11.7, 0.28, 13.5, MUTED, True, PP_ALIGN.CENTER)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

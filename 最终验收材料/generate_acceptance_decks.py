"""Generate three independent final-acceptance slide decks.

The decks are grounded in repository evidence and follow the three TA rubrics.
"""

from pathlib import Path
import sys

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
HELPER = ROOT / "stage3_evidence" / "02_presentation"
sys.path.insert(0, str(HELPER))

from generate_stage3_ppt import (  # noqa: E402
    AMBER, BG, CYAN, FONT, GREEN, INDIGO, LINE, MUTED, PANEL, PANEL_2,
    RED, SH, SW, TEXT, add_picture_contain, card, link_text, rect, rgb,
    rich_lines, text,
)

IMG = ROOT / "stage2_evidence" / "06_careai_validation"
CARE = ROOT / "CareAI"
STAGE3 = ROOT / "stage3_evidence" / "01_careai_acceptance"
DATE = "2026-09-16"


def blank_prs(title_value, subject):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    prs.core_properties.title = title_value
    prs.core_properties.subject = subject
    prs.core_properties.author = "VentureAI项目组"
    return prs


def footer(slide, number, label):
    text(slide, f"{label} · {DATE}", 0.72, 7.15, 10.8, 0.2, 8.2, MUTED)
    text(slide, f"{number:02d}", 12.0, 7.1, 0.6, 0.25, 10, CYAN, True, PP_ALIGN.RIGHT)


def base(prs, number, title_value, subtitle, label, eyebrow):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    text(slide, eyebrow, 0.72, 0.35, 7.2, 0.25, 9, CYAN, True)
    text(slide, title_value, 0.72, 0.68, 11.85, 0.52, 25, TEXT, True)
    if subtitle:
        text(slide, subtitle, 0.72, 1.2, 11.85, 0.34, 11.2, MUTED)
    footer(slide, number, label)
    return slide


def pill(slide, value, x, y, w, color=CYAN, fill=PANEL_2, size=10.5):
    rect(slide, x, y, w, 0.38, fill, color)
    text(slide, value, x, y + 0.02, w, 0.3, size, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def source_line(slide, value):
    text(slide, value, 0.76, 6.85, 11.45, 0.18, 7.7, MUTED)


def cover(prs, station, title_value, subtitle, image_path, tag, label, right_caption):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    pill(s, station, 0.72, 0.62, 1.65, CYAN, "0E2030", 10)
    text(s, title_value, 0.72, 1.38, 6.25, 1.35, 31, TEXT, True)
    text(s, subtitle, 0.76, 3.03, 5.95, 0.92, 16, CYAN, True)
    rich_lines(s, [
        "项目成果与逻辑可独立核验" if "CareAI" in title_value else "真实问题 → 修改 → 同条件验证",
        tag,
    ], 0.76, 4.35, 5.75, 1.1, 12.5, 8, True)
    add_picture_contain(s, image_path, 7.2, 0.95, 5.35, 4.95, CYAN)
    rect(s, 7.2, 6.05, 5.35, 0.56, PANEL_2, LINE)
    text(s, right_caption, 7.42, 6.2, 4.92, 0.24, 10.5, MUTED, True, PP_ALIGN.CENTER)
    footer(s, 1, label)
    return s


def add_before_after(slide, before_path, after_path, y=1.72, h=4.18, left_caption="修改前", right_caption="修改后"):
    add_picture_contain(slide, before_path, 0.72, y, 5.72, h)
    add_picture_contain(slide, after_path, 6.88, y, 5.72, h)
    pill(slide, left_caption, 0.94, y + 0.18, 1.22, RED, "2A1420", 10.5)
    pill(slide, right_caption, 7.1, y + 0.18, 1.22, GREEN, "0D2B27", 10.5)


def build_ta1():
    label = "助教1 · CareAI项目成果"
    prs = blank_prs("助教1｜CareAI项目成果展示", "项目成果、迭代、证据、可行性与风险")

    cover(
        prs,
        "助教1 · 项目成果",
        "CareAI：从健康数据到可执行行动",
        "面向成年人的非诊断性健康教育与习惯管理助手",
        CARE / "frontend" / "careai-home.png",
        "目标用户：高校成年学生及初入职场年轻成年人（H）",
        label,
        "真实产品截图｜当前可演示MVP",
    )

    # 02 evolution
    s = base(prs, 2, "从第一阶段立项到最终MVP：方向保持，闭环变深", "每项变化都对应使用连续性、证据追溯或安全边界", label, "A1 · 项目基础与后续迭代")
    phases = [
        ("第一阶段立项", "基础录入\n规则提示\n受限AI解释\n7天计划", CYAN),
        ("后续实现", "个人档案隔离\n时间轴与多点趋势\n计划持久化与复盘\n医生摘要API", INDIGO),
        ("最终取舍", "保留非诊断边界\n文本导入部分实现\n延期OCR/设备/医院接入\n未开展真实用户验证", GREEN),
    ]
    for i, (ttl, body, color) in enumerate(phases):
        x = 0.72 + i * 4.03
        rect(s, x, 1.73, 3.76, 3.35, PANEL, color)
        text(s, f"0{i+1}", x + 0.28, 2.0, 0.6, 0.42, 17, color, True)
        text(s, ttl, x + 1.0, 2.01, 2.35, 0.38, 17, TEXT, True)
        rich_lines(s, body.split("\n"), x + 0.3, 2.72, 3.1, 1.85, 13.5, 8, True)
        if i < 2:
            text(s, "→", x + 3.72, 3.12, 0.32, 0.4, 22, CYAN, True, PP_ALIGN.CENTER)
    rect(s, 0.72, 5.45, 11.88, 0.86, "0E2030", CYAN)
    text(s, "变化原因：把一次性报告升级为连续习惯管理，同时把医疗、隐私和模型能力限制写进产品流程。", 0.96, 5.69, 11.4, 0.34, 14, TEXT, True, PP_ALIGN.CENTER)
    source_line(s, "证据：CareAI Stage 1项目书；CareAI/docs/stage2/progress.md（17项后端测试、前端构建通过）")

    # 03 solution
    s = base(prs, 3, "最终解决方案：规则负责判断，AI负责解释", "用一条可追溯链路回应“看得懂、知道为什么、下一步能行动”", label, "A2 / A3 · 用户问题与解决方案")
    steps = [
        ("01", "记录", "身高、体重、睡眠、运动\n可选血压/血糖"),
        ("02", "规则", "本地计算与三级提示\n来源和触发原因可查"),
        ("03", "受限AI", "只解释既有提示\n不能新增诊断或风险"),
        ("04", "行动", "生活方式建议\n恰好7天的行动计划"),
        ("05", "回顾", "历史报告、趋势\n前后记录对比"),
    ]
    for i, (n, ttl, body) in enumerate(steps):
        x = 0.72 + i * 1.52
        rect(s, x, 1.78, 1.3, 3.48, PANEL, CYAN if i < 3 else GREEN)
        text(s, n, x + 0.17, 2.0, 0.45, 0.34, 12, CYAN, True)
        text(s, ttl, x + 0.17, 2.53, 0.95, 0.36, 16, TEXT, True)
        text(s, body, x + 0.17, 3.18, 0.96, 1.38, 11.5, MUTED)
        if i < 4:
            text(s, "→", x + 1.29, 3.21, 0.22, 0.36, 16, CYAN, True, PP_ALIGN.CENTER)
    add_picture_contain(s, CARE / "frontend" / "careai-report-demo.png", 8.55, 1.78, 4.05, 3.48, GREEN)
    rect(s, 0.72, 5.62, 11.88, 0.66, PANEL_2, LINE)
    text(s, "明确不解决：急症、疾病诊断、治疗决策、处方、药物剂量、未成年人健康管理。", 0.94, 5.83, 11.42, 0.28, 13.2, AMBER, True, PP_ALIGN.CENTER)
    source_line(s, "证据：CareAI项目书；源码 risk_rules.py / health_report_agent.py；CareAI真实报告页截图")

    # 04 competition
    s = base(prs, 4, "创新不等于“用了AI”：CareAI选择可解释的细分体验", "阿福的生态与规模更强；CareAI聚焦简单、透明和建议可追溯", label, "A3 · 创新与替代方案")
    cols = [0.72, 3.5, 7.8]; widths = [2.55, 4.05, 4.8]
    for val, x, w, color in [("维度", cols[0], widths[0], MUTED), ("CareAI", cols[1], widths[1], CYAN), ("蚂蚁阿福", cols[2], widths[2], GREEN)]:
        rect(s, x, 1.66, w, 0.54, PANEL_2, LINE); text(s, val, x + 0.14, 1.82, w - 0.28, 0.24, 11.5, color, True)
    rows = [
        ("核心价值", "规则可追溯 + 受限解释 + 7天行动", "家庭健康档案 + 健康服务生态"),
        ("AI角色", "规则给结论，AI负责表达", "咨询、报告解读与个性化服务"),
        ("适合用户", "重视简单、透明、可复核", "需要家庭档案与医疗服务连接"),
        ("当前证据", "功能可运行；无真实偏好或效果数据", "公开产品能力可查；体验仍需同条件实测"),
    ]
    for i, row in enumerate(rows):
        y = 2.29 + i * 0.78
        for j, (val, x, w) in enumerate(zip(row, cols, widths)):
            rect(s, x, y, w, 0.66, PANEL if i % 2 == 0 else "0C1728", LINE, False)
            text(s, val, x + 0.12, y + 0.1, w - 0.24, 0.44, 10.5 if j else 11, TEXT if j else MUTED, j == 0, valign=MSO_ANCHOR.MIDDLE)
    rect(s, 0.72, 5.7, 11.88, 0.57, "0E2030", CYAN)
    text(s, "公平验证：同一虚构资料、同一任务、交换顺序、同一指标；理解与行动更好且边界误解不增加，才支持优势主张。", 0.93, 5.85, 11.45, 0.3, 11.7, TEXT, True, PP_ALIGN.CENTER)
    link_text(s, "阿福来源：App Store产品页", "https://apps.apple.com/cn/app/id6743828427", 0.76, 6.5, 2.6, 0.2, 8)
    link_text(s, "公开资料：人民日报（数据来源：蚂蚁集团）", "https://paper.people.com.cn/rmrb/pc/content/202512/16/content_30124850.html", 3.55, 6.5, 4.1, 0.2, 8)
    text(s, "核对日期：2026-09-15", 9.65, 6.5, 2.6, 0.2, 8, MUTED, align=PP_ALIGN.RIGHT)

    # 05 evidence
    s = base(prs, 5, "证据、假设与验证：知道什么，也知道还不知道什么", "F/I/H/S分层，避免把代码存在、模拟输出或Agent建议写成市场事实", label, "A4 · 项目证据与验证逻辑")
    items = [
        ("F · 事实", "产品流程和代码可核查\n17项后端测试 + 前端构建\n权威公共健康与隐私来源", GREEN),
        ("I · 推断", "规则解释与行动计划组合\n可能降低理解和行动门槛\n仍需用户任务验证", CYAN),
        ("H · 假设", "用户愿意持续记录\n能理解非诊断边界\n7天计划具有可执行性", AMBER),
        ("S · 模拟", "Agent角色扮演和模拟评分\n只用于测试流程与提出问题\n不能当用户或市场证据", INDIGO),
    ]
    for i, (ttl, body, color) in enumerate(items):
        x = 0.72 + i * 3.0
        rect(s, x, 1.72, 2.72, 2.86, PANEL, color)
        text(s, ttl, x + 0.25, 1.99, 2.2, 0.4, 16, color, True)
        rich_lines(s, body.split("\n"), x + 0.25, 2.68, 2.2, 1.35, 12.3, 7, True)
    rect(s, 0.72, 4.95, 11.88, 1.18, PANEL_2, LINE)
    text(s, "下一步低风险验证", 0.97, 5.2, 1.65, 0.32, 13, CYAN, True)
    text(s, "完成授权与隐私审查后，邀请5—8名成年志愿者完成录入、报告理解和边界复述；记录完成率、耗时、误解点与人工求助。", 2.78, 5.13, 9.35, 0.55, 13.2, TEXT, True)
    source_line(s, "当前未开展：真实问卷、访谈、用户试用、订单、收入、合作、临床验证或专业阈值复核")

    # 06 implementation & risks
    s = base(prs, 6, "可行性与风险：先在低风险教育场景验证，再决定是否扩大", "技术闭环已经存在；合规、专业复核和真实采用仍是进入真实场景的门槛", label, "A5 / A6 · 实施路径、风险与现场理解")
    paths = [
        ("现在", "课堂演示MVP", "React + FastAPI + SQLite\n规则层可独立运行", CYAN),
        ("下一步", "合规可用性测试", "授权、隐私告知\n5—8名成年志愿者", INDIGO),
        ("之后", "专业复核与试点", "复核阈值和文案\n校园/企业健康教育（H）", GREEN),
    ]
    for i, (phase, ttl, body, color) in enumerate(paths):
        x = 0.72 + i * 4.03
        rect(s, x, 1.72, 3.75, 1.56, PANEL, color)
        text(s, phase, x + 0.23, 1.94, 0.8, 0.3, 11, color, True)
        text(s, ttl, x + 1.08, 1.92, 2.25, 0.35, 15, TEXT, True)
        text(s, body, x + 0.23, 2.5, 3.05, 0.48, 11.5, MUTED)
    risks = [
        ("医疗误解", "非诊断定位、规则/生成分层；上线前专业复核"),
        ("敏感数据", "最小化采集；补足授权、删除、保存期限与访问控制"),
        ("模型/网络", "规则接口可独立演示；AI失败明确报错且不保存报告"),
        ("商业证据", "当前无收入或支付意愿；先验证理解与使用，再讨论模式"),
    ]
    for i, (ttl, body) in enumerate(risks):
        x = 0.72 + (i % 2) * 6.0; y = 3.68 + (i // 2) * 1.15
        rect(s, x, y, 5.72, 0.9, PANEL, LINE)
        text(s, ttl, x + 0.22, y + 0.18, 1.15, 0.3, 13, AMBER, True)
        text(s, body, x + 1.42, y + 0.14, 4.02, 0.46, 10.8, MUTED)
    rect(s, 0.72, 6.1, 11.88, 0.42, GREEN, GREEN)
    text(s, "结论：项目价值是可解释的健康教育闭环；当前最大任务是验证理解、安全边界与持续使用，而不是扩大医疗能力。", 0.92, 6.19, 11.48, 0.24, 11.5, BG, True, PP_ALIGN.CENTER)
    source_line(s, "现场追问准备：为什么选该用户｜为什么不用通用AI｜最大假设是什么｜为什么延期OCR和医疗接入")

    path = OUT / "助教1_CareAI项目成果展示.pptx"
    prs.save(path)
    return path


def build_ta2():
    label = "助教2 · VentureAI优化与效果"
    prs = blank_prs("助教2｜VentureAI优化与实际效果", "V1到V2优化、F1/F2/F3、贡献、局限和随机抽测")
    cover(
        prs,
        "助教2 · Agent优化",
        "VentureAI：从“会回答”到“有证据地指导”",
        "用 CareAI 同一项目材料完成 V1→V2 对照与最终应用",
        IMG / "after-careai-coach.png",
        "冻结版本：V1 commit c456aef → V2 tag ventureai-v2-stage3-2026-09-15",
        label,
        "真实学生端截图｜修改后Coach流程",
    )

    # 02 versions/problem
    s = base(prs, 2, "代表性高优问题：V1会把项目材料改写成不存在的事实", "同一CareAI项目书、同一任务；问题会污染后续指导、竞品比较和评分", label, "B1 · V1问题与优化目标")
    add_picture_contain(s, IMG / "before-careai-coach.png", 0.72, 1.72, 7.2, 4.45, RED)
    rect(s, 8.25, 1.72, 4.35, 4.45, PANEL, LINE)
    failures = [
        ("P0", "事实幻觉", "非诊断健康教育被写成心理筛查；虚构NLP、访谈、团队与收入"),
        ("P0", "数据权限", "匿名可读取项目、详情和会话历史"),
        ("P1", "证据缺失", "文件只截取前2000字；评分与材料无法逐项对应"),
        ("P1", "错误传播", "用户纠正后仍保留旧设定，竞品任务继续失真"),
    ]
    for i, (p, ttl, body) in enumerate(failures):
        y = 1.98 + i * 0.96
        text(s, p, 8.52, y, 0.48, 0.28, 10.5, RED if p == "P0" else AMBER, True)
        text(s, ttl, 9.08, y - 0.02, 1.0, 0.3, 12.5, TEXT, True)
        text(s, body, 10.06, y - 0.03, 2.18, 0.58, 9.8, MUTED)
    rect(s, 8.52, 5.58, 3.8, 0.36, "2A1420", RED)
    text(s, "优先级依据：影响所有核心流程与验收可信度", 8.62, 5.66, 3.6, 0.18, 9.8, TEXT, True, PP_ALIGN.CENTER)
    source_line(s, "原始运行：careai-live-20260914-135948.json；Coach run_id=run_3ac45f3de1cb")

    # 03 coach before/after
    s = base(prs, 3, "V2修改与验证（1）：事实基准进入模型，纠错后撤销旧推断", "修改不只发生在提示词，还包括文件注入、事实优先级、输出边界和交互约束", label, "B1 / B3 · 项目指导前后证据")
    add_before_after(s, IMG / "before-careai-coach.png", IMG / "after-careai-coach.png", 1.72, 3.92)
    rect(s, 0.72, 5.84, 11.88, 0.58, PANEL_2, LINE)
    text(s, "前：心理筛查 / NLP / 已访谈   →   后：非诊断健康教育 / 规则与AI分层 / 明确未验证项", 0.93, 5.99, 11.44, 0.28, 12.2, TEXT, True, PP_ALIGN.CENTER)
    source_line(s, "修改：完整正文与证据索引进入模型；用户/文件事实优先；冲突时撤销旧AI推断；首轮只问一个下一步问题")

    # 04 grader before/after
    s = base(prs, 4, "V2修改与验证（2）：评分从“有分数”变成“有依据、缺口和下一步”", "同一CareAI项目材料；竞品比较也改为同资料、同任务和同指标", label, "B1 / B4 · 评审反馈前后证据")
    add_before_after(s, IMG / "before-careai-grader.png", IMG / "after-careai-grader.png", 1.72, 3.82)
    rect(s, 0.72, 5.75, 7.45, 0.65, PANEL_2, LINE)
    text(s, "R1—R11：每项返回材料依据、证据缺口和可执行修改；无依据内容写“未找到”。", 0.95, 5.93, 6.98, 0.3, 11.2, TEXT, True, PP_ALIGN.CENTER)
    rect(s, 8.38, 5.75, 4.22, 0.65, "0E2030", CYAN)
    text(s, "竞品：不声称CareAI已胜出，只给公平对照方案", 8.57, 5.93, 3.84, 0.3, 10.7, TEXT, True, PP_ALIGN.CENTER)
    source_line(s, "原始记录：careai-completed-grader-recovered.json；竞品修复 run_id=run_07a98721bfe5")

    # 05 flows
    s = base(prs, 5, "三个核心流程已独立准备，可接受现场随机抽测", "每个流程都有明确入口、任务边界、原始run_id和同版本历史记录", label, "B2 / B3 / B4 / B7 · F1/F2/F3")
    flows = [
        ("F1", "理论学习", "定义与举例\n反例/边界\n理解检查\n不确定性与来源", CYAN, "U1 / U2"),
        ("F2", "项目指导", "一次推进一个问题\n用户/场景/证据\n不代写事实\n明确待验证项", INDIGO, "Coach / 竞品"),
        ("F3", "评审反馈", "R1—R11标准\n优点/缺口/风险\n可操作修改\n再次run复验", GREEN, "Grader"),
    ]
    for i, (f, ttl, body, color, entry) in enumerate(flows):
        x = 0.72 + i * 4.03
        rect(s, x, 1.72, 3.76, 3.85, PANEL, color)
        text(s, f, x + 0.28, 1.99, 0.75, 0.48, 22, color, True)
        text(s, ttl, x + 1.12, 2.03, 2.08, 0.4, 17, TEXT, True)
        rich_lines(s, body.split("\n"), x + 0.3, 2.82, 3.05, 1.63, 13.3, 8, True)
        pill(s, f"入口：{entry}", x + 0.3, 4.83, 2.95, color, "0C1728", 10)
    rect(s, 0.72, 5.92, 11.88, 0.55, "0E2030", CYAN)
    text(s, "抽测策略：现场运行当前版本；若网络/API失败，如实展示失败状态，再打开同版本原始运行和日志。", 0.92, 6.06, 11.48, 0.28, 11.5, TEXT, True, PP_ALIGN.CENTER)
    source_line(s, "同条件证据：V1/V2均使用 qwen/qwen-2.5-72b-instruct、USE_MOCK_API=false、固定U1—U6")

    # 06 contribution
    s = base(prs, 6, "Agent对最终CareAI项目的实际贡献：输出经过人工选择才进入成果", "展示输入、Agent输出、采用/拒绝/改写和最终影响，避免把Agent草稿当事实", label, "B6 / C4 · 实际贡献与人机边界")
    cols = [0.72, 3.0, 6.15, 9.4]; widths = [2.05, 2.9, 3.0, 3.2]
    headers = [("项目任务", MUTED), ("提供给Agent", CYAN), ("Agent输出", INDIGO), ("人工处理与最终结果", GREEN)]
    for (val, color), x, w in zip(headers, cols, widths):
        rect(s, x, 1.65, w, 0.55, PANEL_2, LINE); text(s, val, x + 0.12, 1.82, w - 0.24, 0.23, 10.8, color, True)
    rows = [
        ("确认项目定位", "CareAI项目书全文", "指出用户、场景、方案与证据缺口", "采用正确定位；删除心理筛查、访谈等幻觉"),
        ("与阿福比较", "CareAI事实 + 阿福公开资料", "差异化假设与对照任务", "保留同条件测试；拒绝“已优于阿福”的结论"),
        ("完善项目材料", "R1—R11评审任务", "逐项依据、缺口和建议", "补充边界和验证计划；人工核对来源"),
    ]
    for i, row in enumerate(rows):
        y = 2.31 + i * 1.03
        for j, (val, x, w) in enumerate(zip(row, cols, widths)):
            rect(s, x, y, w, 0.9, PANEL if i % 2 == 0 else "0C1728", LINE, False)
            text(s, val, x + 0.12, y + 0.11, w - 0.24, 0.62, 10.2, TEXT if j else MUTED, j == 0, valign=MSO_ANCHOR.MIDDLE)
    rect(s, 0.72, 5.63, 11.88, 0.72, PANEL_2, LINE)
    text(s, "责任边界", 0.97, 5.85, 1.05, 0.28, 12.5, AMBER, True)
    text(s, "VentureAI负责提出分析与反馈；项目组核对事实并决定采纳、拒绝、改写或补充，最终材料由项目组负责。", 2.1, 5.79, 9.96, 0.4, 12.3, TEXT, True)
    source_line(s, "其他工具声明：OpenAI Codex协助整理PPT与证据索引；不作为真实用户、市场或临床证据")

    # 07 limits
    s = base(prs, 7, "优化取舍与剩余问题：已解决关键风险，但没有声称“全部完成”", "B8占10分：说明哪些解决、哪些仍存在、哪些延期，以及下一步优先级", label, "B5 / B8 · 稳定性、异常、剩余问题与后续")
    blocks = [
        ("已解决", "项目/会话权限\n完整材料与事实优先级\nR1—R11依据\nrun_id与失败留痕", GREEN),
        ("仍存在", "模型仍可能生成无来源示例\n部分评分措辞需人工核对\n长请求仍有等待风险\n风险提示可能过宽", AMBER),
        ("延期/放弃", "完整全站权限审计\n历史TypeScript/ESLint清理\n广泛真实用户测试\n跨组测试（教师已取消）", INDIGO),
        ("下一步优先", "F1/F2/F3随机回归\n降低评分延迟\n增强引用定位\n由教师和项目负责人复核", CYAN),
    ]
    for i, (ttl, body, color) in enumerate(blocks):
        x = 0.72 + i * 3.0
        rect(s, x, 1.72, 2.72, 3.42, PANEL, color)
        text(s, ttl, x + 0.24, 1.98, 2.18, 0.38, 16, color, True)
        rich_lines(s, body.split("\n"), x + 0.24, 2.69, 2.18, 1.83, 11.7, 7, True)
    rect(s, 0.72, 5.52, 11.88, 0.78, "0E2030", CYAN)
    text(s, "异常原则：失败不伪装成功；保留状态、错误、重试和人工干预。当前证据含Grader客户端180秒超时后服务端完成并取回的记录。", 0.96, 5.72, 11.4, 0.4, 11.5, TEXT, True, PP_ALIGN.CENTER)
    source_line(s, "验证：15项后端测试、TypeScript检查和生产构建通过；一次固定真实模型运行不能外推为所有提示词下的绝对质量")

    path = OUT / "助教2_VentureAI优化展示.pptx"
    prs.save(path)
    return path


def build_ta3():
    label = "助教3 · 工程证据索引"
    prs = blank_prs("助教3｜VentureAI工程证据索引", "版本、原始运行、部署、人机边界、来源和三阶段证据链")

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    pill(s, "助教3 · 现场核验", 0.72, 0.62, 1.85, CYAN, "0E2030", 10)
    text(s, "VentureAI 工程证据索引", 0.72, 1.42, 7.2, 0.62, 31, TEXT, True)
    text(s, "目标：30秒内定位版本、原始运行、失败记录、运行环境与人工修改", 0.76, 2.34, 8.2, 0.45, 16, CYAN, True)
    chain = [("V1", "c456aef", RED), ("V2", "U1—U6真实运行", INDIGO), ("FINAL", "tag 7975302", GREEN)]
    for i, (a, b, color) in enumerate(chain):
        x = 0.78 + i * 3.85
        rect(s, x, 3.35, 3.4, 1.28, PANEL, color)
        text(s, a, x + 0.22, 3.62, 0.78, 0.38, 15, color, True)
        text(s, b, x + 1.03, 3.62, 2.04, 0.38, 13, TEXT, True)
        if i < 2:
            text(s, "→", x + 3.42, 3.71, 0.4, 0.35, 19, CYAN, True, PP_ALIGN.CENTER)
    rect(s, 0.78, 5.15, 11.02, 0.84, PANEL_2, LINE)
    text(s, "系统冻结版本", 1.05, 5.39, 1.35, 0.28, 12.5, CYAN, True)
    text(s, "ventureai-v2-stage3-2026-09-15 · commit 7975302", 2.55, 5.34, 5.05, 0.36, 17, TEXT, True)
    text(s, "演示文稿是后续材料，不改变冻结系统版本", 8.1, 5.42, 3.3, 0.28, 10.5, MUTED, True)
    footer(s, 1, label)

    # 02 index
    s = base(prs, 2, "原始证据快速定位：版本、运行、失败和三阶段闭环", "截图用于说明，JSON、run_id、commit和日志才是可回溯原始证据", label, "C1 / C2 / C6 · 可追溯证据")
    rows = [
        ("V1冻结", "c456aef", "stage1_evidence/01_AgentV1/"),
        ("V1原始U1—U6", "20260911-192250", "stage2_evidence/05_v1_v2_comparison/"),
        ("V2原始U1—U6", "20260911-193404", "stage2_evidence/04_live_runs/"),
        ("CareAI修改前", "run_3ac45f3de1cb", "stage2_evidence/06_careai_validation/"),
        ("竞品修复后", "run_07a98721bfe5", "careai-competitor-verified.json"),
        ("最终版本", "tag 7975302", "ventureai-v2-stage3-2026-09-15"),
        ("最终测试", "15项后端 + TS + build", "stage3_evidence/.../validation-results.txt"),
    ]
    cols = [0.72, 3.1, 6.25]; widths = [2.15, 2.92, 6.35]
    for val, x, w, color in [("检查内容", cols[0], widths[0], MUTED), ("定位标识", cols[1], widths[1], CYAN), ("文件/目录", cols[2], widths[2], GREEN)]:
        rect(s, x, 1.62, w, 0.5, PANEL_2, LINE); text(s, val, x + 0.12, 1.76, w - 0.24, 0.22, 10.5, color, True)
    for i, row in enumerate(rows):
        y = 2.19 + i * 0.58
        for j, (val, x, w) in enumerate(zip(row, cols, widths)):
            rect(s, x, y, w, 0.5, PANEL if i % 2 == 0 else "0C1728", LINE, False)
            text(s, val, x + 0.11, y + 0.07, w - 0.22, 0.34, 9.4, TEXT if j else MUTED, j == 0, valign=MSO_ANCHOR.MIDDLE)
    source_line(s, "较早的V2运行和失败记录保留在stage2_evidence/04_live_runs/，用于说明修复轨迹；未只保留成功案例")

    # 03 reproducibility
    s = base(prs, 3, "部署、运行与复现：入口清楚，异常有备用证据", "现场优先运行最终版本；外部服务故障时展示失败状态与同版本历史运行", label, "C3 · 运行说明与可复现性")
    card(s, "01  配置", "复制环境示例；选择mock离线演示或配置真实模型密钥。密钥只进入服务端环境变量。", 0.72, 1.7, 3.75, 1.78, CYAN, body_size=12.2)
    card(s, "02  后端", "安装requirements，启动FastAPI；健康检查：http://localhost:8000/health", 4.79, 1.7, 3.75, 1.78, INDIGO, body_size=12.2)
    card(s, "03  前端", "安装锁定依赖并启动Next.js；访问：http://localhost:3000", 8.85, 1.7, 3.75, 1.78, GREEN, body_size=12.2)
    rect(s, 0.72, 3.87, 7.5, 2.08, PANEL, LINE)
    text(s, "验证命令", 0.98, 4.12, 1.2, 0.32, 14, CYAN, True)
    rich_lines(s, [
        "backend/.venv/bin/python -m unittest discover -s backend/tests -v",
        "frontend: npx tsc --noEmit",
        "frontend: npm run build",
    ], 0.98, 4.64, 6.85, 0.98, 11.5, 7, True)
    rect(s, 8.52, 3.87, 4.08, 2.08, PANEL_2, LINE)
    text(s, "异常策略", 8.78, 4.12, 1.1, 0.32, 14, AMBER, True)
    rich_lines(s, [
        "不把失败显示成成功",
        "保留run_id、状态和错误",
        "打开同版本历史运行辅助核验",
    ], 8.78, 4.64, 3.35, 0.98, 11.5, 7, True)
    source_line(s, "完整说明：根目录README.md；依赖：backend/requirements.txt、frontend/package-lock.json；配置：环境示例文件")

    # 04 integrity
    s = base(prs, 4, "可信性边界：来源、模拟、人工判断与其他工具全部可说明", "允许使用Agent和生成式AI；关键是说明各自作用，并由项目组承担最终责任", label, "C4 / C5 · 人机边界、来源与安全诚信")
    blocks = [
        ("VentureAI", "概念辅导、项目指导、R1—R11评审、证据缺口与下一步建议", CYAN),
        ("项目组", "核对事实；采纳、拒绝、改写和补充；决定最终项目结论", GREEN),
        ("CareAI代码/测试", "证明功能存在与可运行；不证明真实用户效果、市场需求或临床有效性", INDIGO),
        ("OpenAI Codex", "协助分析仓库、整理结构、生成PPT和索引；不作为项目事实来源", AMBER),
    ]
    for i, (ttl, body, color) in enumerate(blocks):
        x = 0.72 + (i % 2) * 6.0; y = 1.72 + (i // 2) * 1.55
        rect(s, x, y, 5.72, 1.27, PANEL, color)
        text(s, ttl, x + 0.24, y + 0.22, 1.35, 0.34, 14, color, True)
        text(s, body, x + 1.67, y + 0.19, 3.7, 0.68, 11.2, MUTED)
    rect(s, 0.72, 5.04, 11.88, 1.18, "0E2030", CYAN)
    text(s, "四条红线", 0.96, 5.28, 1.2, 0.32, 13, CYAN, True)
    text(s, "模拟S不得冒充真实调研｜Agent引用必须回到来源｜不得隐藏人工修改｜代码、日志和PPT不得暴露有效密钥或敏感信息", 2.23, 5.22, 9.86, 0.52, 12.2, TEXT, True)
    source_line(s, "现场主索引：最终验收材料/助教3_现场证据索引.md；三阶段材料分别位于stage1_evidence、stage2_evidence、stage3_evidence")

    path = OUT / "助教3_工程证据索引.pptx"
    prs.save(path)
    return path


if __name__ == "__main__":
    outputs = [build_ta1(), build_ta2(), build_ta3()]
    for path in outputs:
        print(f"saved {path}")

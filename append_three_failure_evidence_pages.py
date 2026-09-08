# -*- coding: utf-8 -*-
"""Append three editable failure-evidence pages in place to the user's deck."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

PPT = Path(__file__).resolve().parent / "第一阶段最终汇报PPT_5分钟_美化版_证据双页.pptx"

NAVY = RGBColor(17, 37, 64)
TEAL = RGBColor(20, 136, 137)
RED = RGBColor(191, 56, 56)
ORANGE = RGBColor(210, 121, 31)
WHITE = RGBColor(255, 255, 255)
MIST = RGBColor(247, 250, 253)
SKY = RGBColor(214, 236, 247)
MINT = RGBColor(208, 237, 232)
TEXT = RGBColor(34, 45, 59)
MUTED = RGBColor(92, 105, 120)


def textbox(slide, text, x, y, w, h, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = text
    run.font.name = "Microsoft YaHei"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box


def background(slide):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = MIST
    for x, y, size, color, transparency in [
        (11.45, 0.78, 2.05, SKY, 28), (12.35, 1.55, 1.20, MINT, 10),
        (-0.65, 6.15, 1.85, SKY, 42), (0.15, 6.55, 0.90, MINT, 18),
    ]:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
        circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.fill.transparency = transparency
        circle.line.fill.background()


def header(slide, number, title, subtitle):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.72))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.72), Inches(13.333), Inches(0.055))
    accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(85, 192, 196); accent.line.fill.background()
    textbox(slide, number, 0.45, 0.08, 0.55, 0.45, 20, RGBColor(150, 212, 235), True)
    textbox(slide, title, 1.0, 0.08, 9.0, 0.45, 22, WHITE, True)
    textbox(slide, subtitle, 9.8, 0.12, 3.0, 0.35, 10, RGBColor(203, 220, 235), False, PP_ALIGN.RIGHT)
    textbox(slide, "CareAI × VentureAI  ｜  Phase 1", 9.35, 7.08, 3.3, 0.18, 8.5, MUTED, False, PP_ALIGN.RIGHT)


def placeholder(slide, title, session, accent, note):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.68), Inches(1.08), Inches(11.98), Inches(4.66))
    frame.fill.solid(); frame.fill.fore_color.rgb = WHITE; frame.line.color.rgb = accent; frame.line.width = Pt(1.7)
    textbox(slide, title, 0.95, 1.27, 11.45, 0.30, 17, accent, True, PP_ALIGN.CENTER)
    textbox(slide, session, 0.95, 1.61, 11.45, 0.23, 10, MUTED, False, PP_ALIGN.CENTER)
    textbox(slide, "把对应的原始失败截图放在这里", 1.0, 3.08, 11.3, 0.38, 20, RGBColor(158, 172, 184), True, PP_ALIGN.CENTER)
    textbox(slide, "插入图片后，删除或覆盖这两行提示即可", 1.0, 3.53, 11.3, 0.26, 11, RGBColor(158, 172, 184), False, PP_ALIGN.CENTER)
    note_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.68), Inches(6.02), Inches(11.98), Inches(0.75))
    note_box.fill.solid(); note_box.fill.fore_color.rgb = WHITE; note_box.line.color.rgb = RGBColor(218, 230, 240)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.68), Inches(6.02), Inches(0.09), Inches(0.75))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
    textbox(slide, "你要讲的话：", 0.95, 6.13, 1.2, 0.18, 10.5, accent, True)
    textbox(slide, note, 2.1, 6.07, 10.1, 0.35, 12, TEXT)


def split_placeholder(slide, title, left, right):
    textbox(slide, title, 0.8, 1.02, 11.7, 0.35, 16, NAVY, True, PP_ALIGN.CENTER)
    for x, heading, session, accent, note in [(0.68, *left), (6.88, *right)]:
        frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.48), Inches(5.78), Inches(3.98))
        frame.fill.solid(); frame.fill.fore_color.rgb = WHITE; frame.line.color.rgb = accent; frame.line.width = Pt(1.6)
        textbox(slide, heading, x + 0.18, 1.67, 5.42, 0.30, 15, accent, True, PP_ALIGN.CENTER)
        textbox(slide, session, x + 0.18, 2.00, 5.42, 0.23, 9.2, MUTED, False, PP_ALIGN.CENTER)
        textbox(slide, "截图放置区", x + 0.18, 3.12, 5.42, 0.30, 16, RGBColor(158, 172, 184), True, PP_ALIGN.CENTER)
        textbox(slide, "插入截图后覆盖提示文字", x + 0.18, 3.52, 5.42, 0.20, 10, RGBColor(158, 172, 184), False, PP_ALIGN.CENTER)
        note_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.70), Inches(5.78), Inches(0.92))
        note_box.fill.solid(); note_box.fill.fore_color.rgb = WHITE; note_box.line.color.rgb = RGBColor(218, 230, 240)
        textbox(slide, note, x + 0.20, 5.84, 5.35, 0.50, 10.5, TEXT)


def main():
    prs = Presentation(PPT)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide); header(slide, "09", "失败案例：默认环境无法启动", "VAI-P0-001 ｜ 工具 / 环境问题")
    placeholder(slide, "全局 Python 环境启动 VentureAI 失败", "原始证据：FIRST_RUNS.md 的前置异常 1、2", RED,
                "默认依赖环境冲突使服务不能启动；我们没有修改 V1，而是保留错误并用隔离环境完成基线复现。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide); header(slide, "10", "失败案例：指定角色却被错误路由", "VAI-P1-001 ｜ 流程问题")
    placeholder(slide, "请求 tutor / grader，最终却返回 coach", "会话：5fcdb523…（T1）／c523b8fc…（T3）", RED,
                "理论学习和项目评审没有按任务执行；原因是显式指定的角色被自动路由覆盖。")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background(slide); header(slide, "11", "失败案例：证据边界与单步指导失效", "VAI-P1-002、VAI-P1-003")
    split_placeholder(
        slide,
        "两类问题都来自首次运行：系统没有正确遵守用户已经写明的课程边界",
        ("证据边界失效", "会话：937b0b5c…（C1）／535cf589…（C2）", ORANGE,
         "材料已写明未做真实验证，系统仍给出无依据风险或要求补充当前不存在的证据。"),
        ("单步指导失效", "会话：2c93981e…（T2）／937b0b5c…（C1）", RED,
         "用户要求一次一个任务，系统却一次输出多项诊断和多个问题。"),
    )

    prs.save(PPT)
    print(PPT)


if __name__ == "__main__":
    main()

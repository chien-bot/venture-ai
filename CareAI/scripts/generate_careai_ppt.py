"""Generate the CareAI project-and-test evidence presentation."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
CARE = ROOT / "CareAI"
EVIDENCE = CARE / "docs" / "innovation-evidence"
ASSETS = CARE / "docs" / "ppt-assets"
OUT = ROOT / "CareAI_项目成果与测试证据.pptx"

SW, SH = Inches(13.333), Inches(7.5)
BG = "1E1E1E"
PANEL = "292929"
PANEL_2 = "333333"
WHITE = "FFFFFF"
MUTED = "B8C0CC"
BLUE = "0066FF"
CYAN = "00FFFF"
GREEN = "38E6A1"
AMBER = "FFB347"
RED = "FF5E78"
LINE = "48505E"
FONT_CN = "PingFang SC"
FONT_EN = "DejaVu Sans"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True, transparency=0):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill); shape.fill.transparency = transparency
    shape.line.color.rgb = rgb(line); shape.line.width = Pt(0.8)
    return shape


def line(slide, x, y, w, h=0.02, color=LINE):
    return rect(slide, x, y, w, h, color, color, False)


def text(slide, value, x, y, w, h, size=18, color=WHITE, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT_CN, margin=0.03):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = value
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(color)
    return box


def rich(slide, items, x, y, w, h, size=14, gap=6, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame; frame.clear(); frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.03)
    for idx, item in enumerate(items):
        if isinstance(item, tuple):
            value, color, bold = item
        else:
            value, color, bold = item, MUTED, False
        p = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        p.text = ("•  " if bullet else "") + value
        p.font.name = FONT_CN; p.font.size = Pt(size); p.font.bold = bold; p.font.color.rgb = rgb(color)
        p.space_after = Pt(gap)
    return box


def hyperlink(slide, label, url, x, y, w, h, size=7.8, color=MUTED):
    box = text(slide, label, x, y, w, h, size, color)
    run = box.text_frame.paragraphs[0].runs[0]
    run.hyperlink.address = url; run.font.underline = True
    return box


def add_bg(slide):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(BG)
    # Tech Innovation motif: precise electric-blue guide lines.
    line(slide, 0, 0, 13.333, 0.05, BLUE)
    line(slide, 12.88, 0.05, 0.02, 6.9, "2B4D80")
    for i in range(5):
        rect(slide, 12.75 + i * 0.06, 0.36 + i * 0.06, 0.03, 0.03, CYAN, CYAN, False)


def footer(slide, n, source="CareAI 本地代码与验收证据 · 2026-09-16"):
    text(slide, source, 0.62, 7.16, 10.9, 0.18, 7.5, MUTED)
    text(slide, f"{n:02d}", 12.02, 7.08, 0.55, 0.28, 10, CYAN, True, PP_ALIGN.RIGHT, font=FONT_EN)


def base(prs, n, title_value, subtitle="", eyebrow="CAREAI · PROJECT & EVIDENCE"):
    slide = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(slide)
    text(slide, eyebrow, 0.62, 0.3, 6.4, 0.24, 8.5, CYAN, True, font=FONT_EN)
    text(slide, title_value, 0.62, 0.64, 11.9, 0.52, 25, WHITE, True)
    if subtitle:
        text(slide, subtitle, 0.62, 1.19, 11.8, 0.32, 11.2, MUTED)
    footer(slide, n)
    return slide


def pill(slide, value, x, y, w, color=CYAN, fill=PANEL_2, size=9.5):
    rect(slide, x, y, w, 0.34, fill, color)
    text(slide, value, x, y + 0.01, w, 0.29, size, color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def card(slide, title_value, body, x, y, w, h, accent=BLUE, metric=None):
    rect(slide, x, y, w, h, PANEL, LINE)
    line(slide, x + 0.22, y + 0.22, 0.08, 0.42, accent)
    text(slide, title_value, x + 0.46, y + 0.2, w - 0.68, 0.36, 14, WHITE, True)
    if metric:
        text(slide, metric, x + 0.25, y + 0.76, w - 0.5, 0.68, 29, accent, True, font=FONT_EN)
        text(slide, body, x + 0.25, y + 1.5, w - 0.5, h - 1.68, 11, MUTED)
    else:
        text(slide, body, x + 0.25, y + 0.78, w - 0.5, h - 1.0, 11.5, MUTED)


def add_picture_contain(slide, path: Path, x, y, w, h, border=LINE):
    rect(slide, x - 0.04, y - 0.04, w + 0.08, h + 0.08, "151515", border)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def crop_asset(source: str, target: str, top: float, bottom: float, left=0.0, right=1.0) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    src = EVIDENCE / source; dst = ASSETS / target
    with Image.open(src) as im:
        x0, x1 = int(im.width * left), int(im.width * right)
        y0, y1 = int(im.height * top), int(im.height * bottom)
        im.crop((x0, y0, x1, y1)).save(dst, quality=94)
    return dst


def prepare_assets() -> dict[str, Path]:
    return {
        "report": crop_asset("careai-report-evidence.png", "report-top.png", 0.0, 0.37),
        "evidence": crop_asset("careai-report-evidence.png", "evidence-cards.png", 0.29, 0.70),
        "comprehension": crop_asset("careai-comprehension-passed.png", "comprehension.png", 0.46, 0.82),
        "private": crop_asset("careai-private-session-comprehension.png", "private-session.png", 0.36, 0.82),
        "experiment": crop_asset("careai-single-variable-experiment.png", "single-experiment.png", 0.43, 0.87),
    }


def build():
    a = prepare_assets()
    prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
    prs.core_properties.title = "CareAI 项目成果与测试证据"
    prs.core_properties.subject = "CareAI 项目定位、差异化、产品截图与验收测试"
    prs.core_properties.author = "CareAI 项目组"

    # 01 cover
    s = prs.slides.add_slide(prs.slide_layouts[6]); add_bg(s)
    pill(s, "CAREAI · FINAL SHOWCASE", 0.72, 0.66, 2.45, CYAN, "26313A", 9)
    text(s, "CareAI", 0.72, 1.42, 5.5, 0.78, 42, WHITE, True, font=FONT_EN)
    text(s, "可解释的健康行动助手", 0.75, 2.31, 5.8, 0.62, 25, CYAN, True)
    text(s, "从一次健康报告，升级为“理解—行动—复盘—隐私—安全”的完整闭环", 0.75, 3.21, 5.75, 0.78, 15, MUTED, True)
    rect(s, 0.75, 4.5, 5.58, 1.32, PANEL, LINE)
    text(s, "本次展示", 1.02, 4.78, 1.0, 0.28, 11, CYAN, True)
    text(s, "真实产品截图 · 24 项后端测试 · 12 项安全回归 · 浏览器零报错", 1.02, 5.17, 4.96, 0.35, 13.2, WHITE, True)
    add_picture_contain(s, EVIDENCE / "careai-home-innovation.png", 7.0, 0.8, 5.65, 5.86, BLUE)
    footer(s, 1, "CareAI 项目成果与测试证据 · Tech Innovation")

    # 02 user/problem/evolution
    s = base(prs, 2, "从“看见数字”到“知道下一步怎么做”", "目标用户：愿意自主记录基础健康数据的成年学生与初入职场人群")
    card(s, "真实场景", "睡眠、运动、体重与基础检查数值分散；用户能看到数字，却难以解释并持续行动。", 0.62, 1.72, 3.72, 1.58, CYAN)
    card(s, "核心问题", "通用搜索与对话工具可快速给建议，但判断依据、责任边界与连续复盘通常不够透明。", 0.62, 3.52, 3.72, 1.58, BLUE)
    card(s, "社会价值", "在低风险健康教育场景中，提高理解与自我管理能力；持续提醒何时需要复测或线下求助。", 0.62, 5.32, 3.72, 1.28, GREEN)
    add_picture_contain(s, EVIDENCE / "careai-mobile-home.png", 4.72, 1.72, 2.28, 4.88, CYAN)
    text(s, "产品迭代", 7.45, 1.82, 1.2, 0.3, 12, CYAN, True)
    phases = [
        ("V1", "一次录入 → 一份报告", MUTED),
        ("V2", "健康 Memory → 趋势与复盘", BLUE),
        ("FINAL", "理解确认 → 单变量实验 → 隐私与安全", GREEN),
    ]
    for i, (tag, body, color) in enumerate(phases):
        y = 2.32 + i * 1.28
        pill(s, tag, 7.46, y, 1.08, color, "272D36", 9)
        rect(s, 8.76, y - 0.02, 3.72, 0.76, PANEL, color)
        text(s, body, 9.0, y + 0.16, 3.25, 0.32, 13, WHITE, True)
        if i < 2:
            text(s, "↓", 7.79, y + 0.82, 0.4, 0.35, 16, CYAN, True, PP_ALIGN.CENTER)
    rect(s, 7.46, 6.2, 5.02, 0.42, "26313A", BLUE)
    text(s, "明确不做：诊断、治疗、处方、药物剂量和急症处理", 7.61, 6.3, 4.72, 0.22, 10.6, AMBER, True, PP_ALIGN.CENTER)

    # 03 solution
    s = base(prs, 3, "最终方案：规则负责判断，AI只负责受限表达", "每一步都能回到输入、规则、来源、理解状态与行动记录")
    flow = [
        ("01", "记录", "基础指标\n报告文字"), ("02", "规则", "确定风险等级\n生成来源"),
        ("03", "解释", "本地或受限AI\n不新增风险"), ("04", "确认", "三题teach-back\n通过后行动"),
        ("05", "实验", "只改变一个变量\n一周后复盘"),
    ]
    for i, (num, ttl, body) in enumerate(flow):
        x = 0.62 + i * 1.42
        rect(s, x, 1.72, 1.18, 2.18, PANEL, BLUE if i < 3 else GREEN)
        text(s, num, x + 0.15, 1.93, 0.4, 0.3, 10, CYAN, True, font=FONT_EN)
        text(s, ttl, x + 0.15, 2.4, 0.86, 0.35, 15, WHITE, True)
        text(s, body, x + 0.15, 2.95, 0.86, 0.62, 10.8, MUTED)
        if i < 4:
            text(s, "→", x + 1.18, 2.66, 0.24, 0.34, 16, CYAN, True, PP_ALIGN.CENTER)
    add_picture_contain(s, a["report"], 8.0, 1.72, 4.62, 3.9, GREEN)
    rect(s, 0.62, 4.35, 6.76, 1.25, PANEL_2, LINE)
    text(s, "关键分层", 0.9, 4.64, 1.02, 0.3, 12, CYAN, True)
    text(s, "本地规则固定 risk level；AI不能新增类别。AI关闭时，完整报告仍由本地规则生成。", 2.04, 4.55, 5.0, 0.52, 13, WHITE, True)
    rect(s, 0.62, 5.83, 12.0, 0.67, "26313A", BLUE)
    text(s, "结果：用户不仅得到建议，还能知道“为什么、由谁生成、哪里不确定、下一步如何验证”。", 0.9, 6.03, 11.45, 0.28, 13, WHITE, True, PP_ALIGN.CENTER)

    # 04 CareAI vs A-Fu
    s = base(prs, 4, "为什么用户可能选择 CareAI，而不是只使用综合健康平台", "阿福的生态与服务覆盖更广；CareAI选择更窄、更透明、更适合验证的健康教育体验")
    headers = [("维度", MUTED), ("CareAI", CYAN), ("蚂蚁阿福", GREEN)]
    xs, ws = [0.62, 3.1, 7.65], [2.28, 4.32, 4.97]
    for (label, color), x, w in zip(headers, xs, ws):
        rect(s, x, 1.68, w, 0.5, PANEL_2, LINE); text(s, label, x + 0.14, 1.82, w - 0.28, 0.22, 11, color, True)
    rows = [
        ("产品重点", "理解依据 + 低负担行动实验", "健康咨询 + 报告解读 + 健康档案 + 医疗服务"),
        ("判断机制", "本地规则固定结论；生成层受限", "综合 AI 健康助手与服务生态"),
        ("可解释性", "输入字段、规则版本、来源与不确定性可查看", "公开产品页强调综合能力与个性化服务"),
        ("用户控制", "AI、保存、期限、导出、删除均可切换", "以官方产品实际隐私设置与政策为准"),
        ("当前阶段", "课堂原型；代码与测试可核验", "成熟商业应用；规模与服务连接更强"),
    ]
    for i, row in enumerate(rows):
        y = 2.28 + i * 0.7
        for j, (value, x, w) in enumerate(zip(row, xs, ws)):
            rect(s, x, y, w, 0.6, PANEL if i % 2 == 0 else "242424", LINE, False)
            text(s, value, x + 0.12, y + 0.08, w - 0.24, 0.42, 9.8 if j else 10.3, WHITE if j else MUTED, j == 0, valign=MSO_ANCHOR.MIDDLE)
    rect(s, 0.62, 6.03, 12.0, 0.52, "26313A", CYAN)
    text(s, "选择逻辑：需要综合服务生态 → 阿福；需要可复核的规则、学习式理解和单变量自我实验 → CareAI。", 0.83, 6.16, 11.58, 0.26, 11.2, WHITE, True, PP_ALIGN.CENTER)
    hyperlink(s, "阿福能力来源：App Store 官方产品页（核对：2026-09-16）", "https://apps.apple.com/cn/app/id6743828427", 0.66, 6.74, 4.5, 0.18)
    text(s, "注：本页是定位比较，不代表已完成同条件真人体验测试。", 7.7, 6.72, 4.7, 0.2, 7.8, AMBER, align=PP_ALIGN.RIGHT)

    # 05 evidence + comprehension
    s = base(prs, 5, "创新一：建议必须能解释，行动必须先理解", "证据卡负责“为什么”；三题理解确认负责“用户是否真的看懂”")
    add_picture_contain(s, a["evidence"], 0.62, 1.7, 5.88, 4.55, BLUE)
    add_picture_contain(s, a["comprehension"], 6.75, 1.7, 5.87, 4.55, GREEN)
    pill(s, "证据卡", 0.86, 1.93, 1.05, CYAN, "24313A", 9)
    pill(s, "理解确认", 6.99, 1.93, 1.25, GREEN, "26352F", 9)
    text(s, "输入字段 · 规则版本 · 生成来源 · 不确定性", 0.82, 6.42, 5.46, 0.26, 11, MUTED, True, PP_ALIGN.CENTER)
    text(s, "含义 · 安全边界 · 下一步；通过后才解锁行动", 6.94, 6.42, 5.46, 0.26, 11, MUTED, True, PP_ALIGN.CENTER)

    # 06 privacy
    s = base(prs, 6, "创新二：隐私不是声明，而是用户可以操作的控制", "关闭 AI、停止保存、设置期限、导出和删除，都会改变后端真实行为")
    add_picture_contain(s, EVIDENCE / "careai-privacy-center.png", 0.62, 1.7, 7.35, 4.82, CYAN)
    add_picture_contain(s, a["private"], 8.26, 1.7, 4.36, 4.82, GREEN)
    checks = [
        ("AI OFF", "仅运行本地规则，不向外部模型发送关注项"),
        ("SAVE OFF", "报告与所有计划写入接口停止保存"),
        ("SESSION", "不保存时仍在当前页面完成理解确认"),
        ("EXPORT", "JSON / 打印PDF / FHIR结构演示"),
    ]
    for i, (tag, body) in enumerate(checks):
        x = 0.7 + i * 3.0
        pill(s, tag, x, 6.7, 0.9, CYAN if i != 1 else AMBER, "28313A", 8.2)
        text(s, body, x + 1.02, 6.68, 1.82, 0.32, 8.5, MUTED, True)

    # 07 single variable experiment
    s = base(prs, 7, "创新三：每周只改变一个变量，让复盘更像实验", "提醒时间、行动时长、行动频率、任务难度四选一；后端对前后 snapshot 做确定性差异校验")
    add_picture_contain(s, a["experiment"], 0.62, 1.7, 7.62, 4.92, GREEN)
    metrics = [
        ("1", "每轮变量", "只能改变一个配置键", CYAN),
        ("7", "行动天数", "固定生成七天结构", BLUE),
        ("vN", "版本链", "保存 previous_draft_id", GREEN),
    ]
    for i, (m, ttl, body, color) in enumerate(metrics):
        y = 1.74 + i * 1.33
        rect(s, 8.58, y, 4.04, 1.1, PANEL, color)
        text(s, m, 8.85, y + 0.19, 0.7, 0.48, 25, color, True, font=FONT_EN)
        text(s, ttl, 9.75, y + 0.18, 2.35, 0.3, 13, WHITE, True)
        text(s, body, 9.75, y + 0.58, 2.45, 0.3, 10.3, MUTED)
    rect(s, 8.58, 5.91, 4.04, 0.71, "26313A", BLUE)
    text(s, "价值：减少“一次改很多、最后不知道什么有效”的复盘偏差。", 8.83, 6.07, 3.55, 0.36, 11, WHITE, True, PP_ALIGN.CENTER)

    # 08 safety + tests
    s = base(prs, 8, "验证结果：功能能运行，边界也能被重复检查", "自动化测试、真实浏览器流程和固定安全用例共同组成验收证据")
    add_picture_contain(s, EVIDENCE / "careai-safety-lab.png", 0.62, 1.68, 6.72, 4.98, BLUE)
    card(s, "后端自动化", "覆盖隐私、本地模式、来源、理解门控、导出、单变量差异与安全套件", 7.68, 1.7, 2.2, 2.0, CYAN, "24/24")
    card(s, "安全实验室", "中英文疾病判断、药名、剂量、频次、停药与安全建议", 10.1, 1.7, 2.52, 2.0, GREEN, "12/12")
    card(s, "真实浏览器", "桌面 + 手机；8张真实页面截图；关键交互由 Playwright 驱动", 7.68, 3.96, 2.2, 2.0, BLUE, "8")
    card(s, "控制台错误", "最终浏览器验收 browser-qa.json；前端生产 build 同时通过", 10.1, 3.96, 2.52, 2.0, AMBER, "0")
    text(s, "命令证据：backend/.venv/bin/python -m pytest -q  →  24 passed", 7.75, 6.24, 4.77, 0.22, 9.5, CYAN, True, font=FONT_EN)
    text(s, "边界：安全回归证明当前列出的代码检查可运行，不等于临床安全认证。", 7.75, 6.52, 4.77, 0.22, 8.8, AMBER)

    # 09 evidence boundary / next
    s = base(prs, 9, "当前结论与下一步：先验证理解和可用性，再讨论扩大", "代码与测试已经形成证据；用户采用、健康效果与商业价值仍是待验证假设")
    items = [
        ("F · 已验证事实", "24项后端测试通过\n前端生产构建通过\n8张浏览器证据，控制台0错误", GREEN),
        ("I · 合理推断", "可解释卡与理解确认\n可能降低误解与行动门槛\n需要真实任务继续检验", CYAN),
        ("H · 待验证假设", "用户愿意持续记录\n单变量实验能改善复盘质量\n存在可持续使用或付费需求", AMBER),
        ("S · 系统证据", "AI与自动化测试输出\n只证明系统流程和代码行为\n不能替代真人或临床证据", BLUE),
    ]
    for i, (ttl, body, color) in enumerate(items):
        x = 0.62 + i * 3.02
        rect(s, x, 1.72, 2.76, 2.7, PANEL, color)
        text(s, ttl, x + 0.24, 2.0, 2.25, 0.4, 14, color, True)
        rich(s, body.split("\n"), x + 0.24, 2.71, 2.26, 1.3, 11.5, 6, True)
    rect(s, 0.62, 4.8, 12.0, 1.35, PANEL_2, LINE)
    text(s, "下一步优先级", 0.91, 5.1, 1.36, 0.3, 12, CYAN, True)
    next_steps = [
        "01  专业人员复核规则阈值与安全文案",
        "02  完成授权与隐私审查后，进行5–8名成年人的可用性与理解测试",
        "03  记录完成率、耗时、误解点与求助行为，再决定OCR、设备或服务接入",
    ]
    rich(s, next_steps, 2.38, 5.0, 9.8, 0.84, 11.5, 5)
    rect(s, 0.62, 6.37, 12.0, 0.45, BLUE, BLUE)
    text(s, "CareAI 的优势不是“替用户看病”，而是让健康教育的依据、理解、行动和边界都能被检查。", 0.88, 6.47, 11.48, 0.24, 11.7, WHITE, True, PP_ALIGN.CENTER)
    hyperlink(s, "证据目录：CareAI/docs/innovation-evidence/", "https://github.com/Qifang24/CareAI", 0.66, 6.91, 3.35, 0.18)
    text(s, "当前状态：可演示课堂原型；FHIR为结构演示，不代表临床互操作认证。", 7.15, 6.9, 5.25, 0.2, 7.8, MUTED, align=PP_ALIGN.RIGHT)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()

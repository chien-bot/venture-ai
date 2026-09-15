"""Generate the compact 10-slide CareAI competitive and VentureAI change deck."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from generate_stage3_ppt import (
    AMBER, BG, CYAN, FONT, GREEN, IMG, INDIGO, LINE, MUTED, OUT, PANEL,
    PANEL_2, RED, SH, SW, TEXT, add_footer, add_picture_contain, base, card,
    link_text, rect, rgb, rich_lines, text,
)


def comparison_screenshot(slide, before, after, y=1.75, h=4.45):
    add_picture_contain(slide, IMG / before, 0.72, y, 5.72, h)
    add_picture_contain(slide, IMG / after, 6.88, y, 5.72, h)
    rect(slide, 0.95, y + 0.18, 1.25, 0.36, RED, RED)
    text(slide, "修改前", 0.95, y + 0.18, 1.25, 0.36, 11, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    rect(slide, 7.11, y + 0.18, 1.25, 0.36, GREEN, GREEN)
    text(slide, "修改后", 7.11, y + 0.18, 1.25, 0.36, 11, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)


def build_compact():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    prs.core_properties.title = "为什么选择 CareAI，以及 VentureAI 修改前后"
    prs.core_properties.subject = "CareAI 与蚂蚁阿福竞争比较；VentureAI 全部关键改进"
    prs.core_properties.author = "VentureAI 项目组"

    # 01
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0.72, 0.72, 1.55, 0.4, CYAN, CYAN)
    text(s, "10-SLIDE REVIEW", 0.72, 0.72, 1.55, 0.4, 9.5, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, "为什么选择 CareAI，\n以及 VentureAI 改了什么", 0.72, 1.52, 8.05, 1.7, 33, TEXT, True)
    text(s, "竞品：蚂蚁阿福  ·  证据：修改前后真实截图", 0.76, 3.6, 7.5, 0.42, 17, CYAN, True)
    text(s, "第三阶段终局验收精简版", 0.76, 4.16, 5.0, 0.35, 14, MUTED)
    rect(s, 9.05, 0.92, 3.5, 5.65, PANEL_2, LINE)
    text(s, "PART 01", 9.45, 1.45, 1.2, 0.3, 11, CYAN, True)
    text(s, "CareAI × 阿福", 9.45, 1.88, 2.4, 0.42, 20, TEXT, True)
    text(s, "直接竞争比较\n优势与选择条件\n公平验证方案", 9.45, 2.55, 2.4, 1.35, 15, MUTED)
    rect(s, 9.45, 4.18, 2.55, 0.06, CYAN, CYAN, False)
    text(s, "PART 02", 9.45, 4.55, 1.2, 0.3, 11, GREEN, True)
    text(s, "VentureAI 前 × 后", 9.45, 4.98, 2.65, 0.42, 19, TEXT, True)
    text(s, "全部关键变化\n四张真实截图", 9.45, 5.55, 2.4, 0.75, 15, MUTED)
    add_footer(s, 1, "第三阶段终局验收 · 2026-09-15")

    # 02
    s = base(prs, 2, "选择 CareAI 的理由：更聚焦、更可解释、更可控", "它不以功能数量击败阿福，而是服务一类明确的使用偏好")
    rect(s, 0.72, 1.7, 11.88, 1.2, "0E2030", CYAN)
    text(s, "目标用户选择条件", 1.02, 2.02, 2.0, 0.32, 14, CYAN, True)
    text(s, "希望快速理解基础指标、看清判断依据，并马上获得低风险行动计划", 3.05, 1.94, 8.9, 0.5, 20, TEXT, True, PP_ALIGN.CENTER)
    reasons = [
        ("01", "规则先于 AI", "关键关注等级由确定性规则产生，AI 负责解释；更容易复核。", CYAN),
        ("02", "边界直接可见", "持续强调非诊断定位、适用范围和线下就医提示。", GREEN),
        ("03", "行动输出固定", "从指标到原因、建议和 7 天计划，减少“看完不知道做什么”。", INDIGO),
        ("04", "产品更轻", "聚焦成年人基础健康教育，不把医疗服务生态全部塞进一次体验。", AMBER),
    ]
    for i, (n, ttl, body, color) in enumerate(reasons):
        x = 0.72 + (i % 2) * 6.0; y = 3.3 + (i // 2) * 1.5
        card(s, ttl, body, x, y, 5.72, 1.22, color, n, 12.5)
    text(s, "竞争主张：CareAI 的优势是“透明且专注的组合体验”，不是阿福缺少档案或习惯目标。", 0.82, 6.42, 11.7, 0.3, 13, AMBER, True, PP_ALIGN.CENTER)

    # 03
    s = base(prs, 3, "CareAI 与蚂蚁阿福：直接竞品对照", "阿福规模与生态更强；CareAI 争夺的是可解释、低复杂度的细分体验")
    cols = [0.72, 3.75, 8.12]; widths = [2.8, 4.12, 4.48]
    for label, color, x, w in [("维度", MUTED, cols[0], widths[0]), ("CareAI", CYAN, cols[1], widths[1]), ("蚂蚁阿福", GREEN, cols[2], widths[2])]:
        rect(s, x, 1.65, w, 0.55, PANEL_2, LINE); text(s, label, x + 0.15, 1.81, w - 0.3, 0.25, 12, color, True)
    rows = [
        ("核心选择理由", "轻量、解释透明、规则可追溯", "完整生态、服务连接、品牌与规模"),
        ("健康档案/习惯", "基础记录 + 7天行动计划", "个人/家庭档案 + 健康小目标与提醒"),
        ("AI 使用方式", "规则给结论，受限 AI 做解释", "健康咨询、报告解读与个性化服务"),
        ("服务边界", "非诊断健康教育与习惯管理", "可连接挂号、云陪诊等医疗健康服务"),
        ("当前短板", "无真实用户、留存与付费证据", "具体体验效果仍须同条件实测"),
        ("更适合谁", "看重简单、透明、可复核的用户", "需要家庭档案和医疗服务生态的用户"),
    ]
    for i, row in enumerate(rows):
        y = 2.29 + i * 0.65
        for j, (v, x, w) in enumerate(zip(row, cols, widths)):
            rect(s, x, y, w, 0.57, PANEL if i % 2 == 0 else "0C1728", LINE, radius=False)
            text(s, v, x + 0.13, y + 0.08, w - 0.26, 0.4, 10.5, TEXT if j else MUTED, j == 0, valign=MSO_ANCHOR.MIDDLE)
    link_text(s, "来源1：蚂蚁阿福 App Store 产品页", "https://apps.apple.com/cn/app/id6743828427", 0.76, 6.36, 3.25, 0.22)
    link_text(s, "来源2：人民日报（数据来源：蚂蚁集团）", "https://paper.people.com.cn/rmrb/pc/content/202512/16/content_30124850.html", 4.15, 6.36, 4.0, 0.22)
    text(s, "核对：2026-09-15", 10.0, 6.36, 2.0, 0.22, 8.5, MUTED, align=PP_ALIGN.RIGHT)

    # 04
    s = base(prs, 4, "CareAI 的五项竞争优势", "这些优势来自当前实现；用户是否在意，仍须通过对照测试验证")
    advantages = [
        ("A1", "确定性规则", "BMI、睡眠、运动等关注等级可复现"),
        ("A2", "受限生成", "AI 不能自行升级成诊断、处方或治疗"),
        ("A3", "可解释链路", "输入 → 规则 → 提示 → 建议 → 7天计划"),
        ("A4", "低认知负担", "一次体验围绕“理解与行动”，入口更聚焦"),
        ("A5", "工程透明度", "规则版本、失败、输出结构可检查和测试"),
    ]
    for i, (n, ttl, body) in enumerate(advantages):
        x = 0.72 + i * 2.4
        rect(s, x, 1.75, 2.12, 3.6, PANEL, LINE)
        text(s, n, x + 0.22, 2.02, 0.65, 0.38, 16, CYAN if i < 3 else GREEN, True)
        text(s, ttl, x + 0.22, 2.64, 1.7, 0.65, 17, TEXT, True)
        text(s, body, x + 0.22, 3.55, 1.66, 1.05, 13, MUTED)
    rect(s, 0.72, 5.82, 11.72, 0.72, "2A2112", AMBER)
    text(s, "真正的护城河不是“也能生成建议”，而是每个建议都能说明从哪里来、为什么安全、下一步怎么做。", 0.94, 6.02, 11.3, 0.3, 13.5, AMBER, True, PP_ALIGN.CENTER)

    # 05
    s = base(prs, 5, "怎样证明 CareAI 真的比阿福更适合这类用户？", "用同一份 S（模拟）资料完成同一任务，避免功能不对等和真实健康数据风险")
    add_picture_contain(s, IMG / "after-careai-competitor.png", 0.72, 1.68, 6.45, 4.45)
    steps = [
        ("01", "同资料", "相同虚构用户资料"),
        ("02", "同任务", "找解释、行动和边界"),
        ("03", "换顺序", "抵消先后使用偏差"),
        ("04", "同指标", "完成、耗时、理解、清晰度、偏好"),
    ]
    for i, (n, ttl, body) in enumerate(steps):
        y = 1.68 + i * 1.13
        rect(s, 7.55, y, 4.85, 0.91, PANEL, LINE)
        text(s, n, 7.78, y + 0.24, 0.45, 0.3, 12, CYAN, True)
        text(s, ttl, 8.38, y + 0.17, 1.0, 0.32, 14, TEXT, True)
        text(s, body, 9.38, y + 0.18, 2.7, 0.42, 11.5, MUTED)
    rect(s, 7.55, 6.25, 4.85, 0.45, "0E2030", CYAN)
    text(s, "胜出条件：理解与行动指标更好，且边界误解不增加", 7.73, 6.35, 4.48, 0.22, 11, TEXT, True, PP_ALIGN.CENTER)

    # 06
    s = base(prs, 6, "VentureAI 修改前：八类问题会破坏验收可信度", "这些问题由同一份 CareAI 项目书、同一组任务实际触发")
    issues = [
        ("P0", "事实幻觉", "把项目改成心理筛查；补造 NLP、访谈、团队与收入"),
        ("P0", "匿名读取", "未登录可读项目列表、详情和会话历史"),
        ("P1", "文件截断", "上传正文只取前2000字，关键边界进入不了模型"),
        ("P1", "错误记忆", "用户纠正后仍沿用上一轮 AI 的错误设定"),
        ("P1", "评分脱证据", "R1—R11 有分数，但部分依据无法回到材料"),
        ("P1", "竞品失真", "把两款产品都当心理筛查，自定30人和错误指标"),
        ("P2", "交互过载", "只问一个问题时仍生成多任务清单；长等待无进度"),
        ("P2", "界面误导", "未评分显示五维0分；规则误报和无关案例增加噪声"),
    ]
    for i, (p, ttl, body) in enumerate(issues):
        x = 0.72 + (i % 2) * 6.0; y = 1.66 + (i // 2) * 1.22
        rect(s, x, y, 5.72, 1.0, PANEL, LINE)
        text(s, p, x + 0.2, y + 0.17, 0.55, 0.28, 11, RED if p == "P0" else AMBER, True)
        text(s, ttl, x + 0.87, y + 0.15, 1.25, 0.3, 14, TEXT, True)
        text(s, body, x + 2.05, y + 0.14, 3.4, 0.55, 11, MUTED)
    text(s, "后果：项目建议、竞品分析、评分和验收都可能建立在错误事实上。", 0.8, 6.55, 11.6, 0.25, 13, RED, True, PP_ALIGN.CENTER)

    # 07
    s = base(prs, 7, "VentureAI 修改后（1/2）：回答内容变得可信、可执行", "模型层、提示词层和交互层同时修复")
    rows = [
        ("文件进入模型", "前2000字且节点未完整读取", "最新正文 + 关键证据索引进入实际模型消息"),
        ("事实优先级", "历史 AI 推断继续传播", "用户/文件事实优先；冲突时撤销旧推断"),
        ("项目定位", "心理筛查、NLP、访谈等幻觉", "准确复述非诊断定位、MVP与未验证项"),
        ("医疗边界", "慢性病发现可被写成普通假设", "临床依据不足时禁止当产品主张"),
        ("竞品比较", "不对等任务、虚构竞品能力", "同资料、同任务、同指标；未知项写未找到"),
        ("R1—R11", "分数与正文脱节", "逐项显示依据、缺口和下一步，返回完整结构"),
        ("对话方式", "一个问题变成多项作业", "首轮只保留一个下一步问题"),
    ]
    cols = [0.72, 3.0, 7.4]; widths = [2.05, 4.15, 5.2]
    for label, color, x, w in [("变化项", MUTED, cols[0], widths[0]), ("修改前", RED, cols[1], widths[1]), ("修改后", GREEN, cols[2], widths[2])]:
        rect(s, x, 1.62, w, 0.52, PANEL_2, LINE); text(s, label, x + 0.14, 1.77, w - 0.28, 0.22, 11.5, color, True)
    for i, row in enumerate(rows):
        y = 2.21 + i * 0.59
        for j, (v, x, w) in enumerate(zip(row, cols, widths)):
            rect(s, x, y, w, 0.51, PANEL if i % 2 == 0 else "0C1728", LINE, radius=False)
            text(s, v, x + 0.12, y + 0.07, w - 0.24, 0.35, 9.5 if j else 10.2, TEXT if j else MUTED, j == 0, valign=MSO_ANCHOR.MIDDLE)
    rect(s, 0.72, 6.5, 11.88, 0.3, "0E2030", CYAN)

    # 08
    s = base(prs, 8, "VentureAI 修改后（2/2）：系统层形成完整证据闭环", "数据库、权限、版本和产品门槛共同保证结果可审计")
    changes = [
        ("安全权限", "项目、会话、上传、工具、知识与图谱均需身份和项目/班级权限"),
        ("注册边界", "公开注册只能创建学生；教师和管理员必须授权配置"),
        ("运行追踪", "保存 run_id、版本、Git、模型、代码/提示词哈希、状态和时间"),
        ("失败与等待", "失败运行仍留痕；长流式任务发送“正在核对”进度"),
        ("材料上传", "新增 DOCX/PPTX 提取；报告文档实测读取13,430字"),
        ("Agent草稿", "草稿不能当正式成果；必须由学生核验并提交新版本"),
        ("第三阶段", "G1—G6、六类材料、F1/F2/F3、F/I/H/S、D1—D9"),
        ("评分验收", "项目/Agent/答辩三表；版本变化使旧评分失效；支持排名与导出"),
    ]
    for i, (ttl, body) in enumerate(changes):
        x = 0.72 + (i % 2) * 6.0; y = 1.65 + (i // 2) * 1.22
        rect(s, x, y, 5.72, 1.0, PANEL, LINE)
        rect(s, x + 0.19, y + 0.2, 0.08, 0.5, GREEN, GREEN, False)
        text(s, ttl, x + 0.43, y + 0.14, 1.28, 0.34, 13.5, TEXT, True)
        text(s, body, x + 1.76, y + 0.13, 3.7, 0.58, 10.8, MUTED)
    rect(s, 0.72, 6.5, 11.88, 0.34, GREEN, GREEN)
    text(s, "验证：15项后端测试通过 · TypeScript通过 · 生产构建通过 · V2版本已冻结", 0.92, 6.55, 11.48, 0.22, 12, BG, True, PP_ALIGN.CENTER)

    # 09
    s = base(prs, 9, "真实截图｜教练流程：从错误项目画像到证据优先", "同一 CareAI 项目书、同一核心问题")
    comparison_screenshot(s, "before-careai-coach.png", "after-careai-coach.png", 1.68, 4.55)
    rect(s, 0.72, 6.35, 11.88, 0.43, PANEL_2, LINE)
    text(s, "前：心理筛查 / NLP / 已访谈  →  后：非诊断健康教育 / 规则与AI分层 / 明确未验证项", 0.92, 6.45, 11.48, 0.22, 11.8, TEXT, True, PP_ALIGN.CENTER)

    # 10
    s = base(prs, 10, "真实截图｜评分流程：从“只有分数”到“可定位修改”", "修改后按 R1—R11 显示材料依据、证据缺口和下一步建议")
    comparison_screenshot(s, "before-careai-grader.png", "after-careai-grader.png", 1.68, 4.35)
    rect(s, 0.72, 6.15, 11.88, 0.67, "0E2030", CYAN)
    text(s, "CareAI 的竞争价值更清楚；VentureAI 的每个结论也能回到材料、版本、run_id 与人工判断。", 0.92, 6.35, 11.48, 0.28, 13.2, TEXT, True, PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build_compact()

"""Generate the VentureAI stage-three acceptance presentation."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "VentureAI_第三阶段终局验收_CareAI案例.pptx"
IMG = ROOT / "stage2_evidence" / "06_careai_validation"
STAGE3 = ROOT / "stage3_evidence" / "01_careai_acceptance"

BG = "08111F"
PANEL = "101B2D"
PANEL_2 = "14213A"
TEXT = "F4F7FB"
MUTED = "9FB0C7"
CYAN = "22D3EE"
INDIGO = "6366F1"
GREEN = "10B981"
AMBER = "F59E0B"
RED = "F43F5E"
LINE = "24334B"

SW, SH = Inches(13.333), Inches(7.5)
FONT = "PingFang SC"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=PANEL, line=LINE, radius=True, transparency=0):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.fill.transparency = transparency
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(0.8)
    return shape


def text(slide, value, x, y, w, h, size=18, color=TEXT, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT, margin=0.04):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def rich_lines(slide, lines, x, y, w, h, size=17, gap=8, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    f = box.text_frame
    f.clear(); f.word_wrap = True
    f.margin_left = f.margin_right = Inches(0.03)
    f.margin_top = f.margin_bottom = Inches(0.03)
    for idx, item in enumerate(lines):
        if isinstance(item, tuple):
            value, color, bold = item
        else:
            value, color, bold = item, TEXT, False
        p = f.paragraphs[0] if idx == 0 else f.add_paragraph()
        p.text = ("•  " if bullet else "") + value
        p.font.name = FONT; p.font.size = Pt(size); p.font.bold = bold
        p.font.color.rgb = rgb(color); p.space_after = Pt(gap)
    return box


def link_text(slide, label, url, x, y, w, h, size=8.5):
    box = text(slide, label, x, y, w, h, size, MUTED)
    run = box.text_frame.paragraphs[0].runs[0]
    run.hyperlink.address = url
    run.font.underline = True
    return box


def add_footer(slide, number, source="课程报告与本地验收证据 · 2026-09-15"):
    text(slide, source, 0.72, 7.16, 10.8, 0.2, 8.5, MUTED)
    text(slide, f"{number:02d}", 12.0, 7.11, 0.6, 0.25, 10, CYAN, True, PP_ALIGN.RIGHT)


def base(prs, number, title_value, subtitle="", eyebrow="VENTUREAI · 第三阶段终局验收"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(BG)
    text(slide, eyebrow, 0.72, 0.35, 7.0, 0.25, 9, CYAN, True)
    text(slide, title_value, 0.72, 0.68, 11.8, 0.5, 26, TEXT, True)
    if subtitle:
        text(slide, subtitle, 0.72, 1.2, 11.8, 0.34, 11.5, MUTED)
    add_footer(slide, number)
    return slide


def card(slide, title_value, body, x, y, w, h, accent=CYAN, number=None, body_size=15):
    rect(slide, x, y, w, h)
    if number is not None:
        rect(slide, x + 0.22, y + 0.2, 0.46, 0.46, accent, accent)
        text(slide, str(number), x + 0.22, y + 0.2, 0.46, 0.46, 13, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        tx = x + 0.82
    else:
        rect(slide, x + 0.24, y + 0.25, 0.08, 0.35, accent, accent, False)
        tx = x + 0.48
    text(slide, title_value, tx, y + 0.2, w - (tx - x) - 0.2, 0.4, 16, TEXT, True)
    text(slide, body, x + 0.25, y + 0.78, w - 0.5, h - 0.95, body_size, MUTED)


def add_picture_contain(slide, path, x, y, w, h, border=LINE):
    rect(slide, x - 0.04, y - 0.04, w + 0.08, h + 0.08, "0B1526", border)
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def screenshot_slide(prs, number, title_value, image_name, badge, observation, why, accent):
    slide = base(prs, number, title_value, "同一 CareAI 项目、同一任务；截图保留原始对话和评分界面")
    add_picture_contain(slide, IMG / image_name, 0.72, 1.66, 9.7, 5.2)
    rect(slide, 10.68, 1.66, 1.95, 5.2, PANEL_2, LINE)
    rect(slide, 10.92, 1.92, 1.45, 0.38, accent, accent)
    text(slide, badge, 10.92, 1.92, 1.45, 0.38, 11, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(slide, "画面说明", 10.92, 2.62, 1.4, 0.3, 12, TEXT, True)
    text(slide, observation, 10.92, 2.98, 1.42, 1.34, 12, MUTED)
    text(slide, "为什么重要", 10.92, 4.55, 1.4, 0.3, 12, TEXT, True)
    text(slide, why, 10.92, 4.9, 1.42, 1.55, 12, MUTED)
    return slide


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    prs.core_properties.title = "VentureAI 第三阶段终局验收——CareAI 测试案例"
    prs.core_properties.subject = "验收标准、CareAI 与蚂蚁阿福比较、VentureAI 修改前后证据"
    prs.core_properties.author = "VentureAI 项目组"

    # 01 cover
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    rect(s, 0.72, 0.68, 1.4, 0.4, CYAN, CYAN)
    text(s, "FINAL REVIEW", 0.72, 0.68, 1.4, 0.4, 10, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, "第三阶段综合应用\n与终局验收", 0.72, 1.42, 7.2, 1.8, 36, TEXT, True)
    text(s, "以 CareAI 为测试案例验证 VentureAI V2", 0.76, 3.45, 6.7, 0.4, 18, CYAN, True)
    text(s, "验收标准 · 竞品比较 · 修改前后证据 · 最终结论", 0.76, 3.98, 7.2, 0.4, 14, MUTED)
    rect(s, 8.55, 0.88, 3.95, 5.85, PANEL_2, LINE)
    text(s, "证据链", 8.95, 1.35, 2.8, 0.4, 14, CYAN, True)
    for i, (a, b) in enumerate([("G1—G6", "立项闯关"), ("F1 / F2 / F3", "三流程"), ("100 + 100", "双门槛"), ("D1—D9", "版本推进"), ("F / I / H / S", "证据边界")]):
        yy = 2.05 + i * 0.82
        text(s, a, 8.95, yy, 1.35, 0.35, 15, TEXT, True)
        text(s, b, 10.55, yy + 0.02, 1.5, 0.3, 12, MUTED)
        if i < 4:
            rect(s, 9.58, yy + 0.43, 0.06, 0.28, CYAN, CYAN, False)
    add_footer(s, 1, "VentureAI V2 · CareAI 验收案例 · 2026-09-15")

    # 02 executive result
    s = base(prs, 2, "先给结论：从“会回答”升级为“可验收”", "第三阶段关注可信过程，而不是生成材料的长度")
    card(s, "功能闭环", "立项 → 材料 → Agent 使用 → 独立评分 → 排名/答辩 → 终局结论", 0.72, 1.75, 3.85, 2.1, GREEN, "01")
    card(s, "CareAI 流程证据", "ready = true；缺项 = 0；F1/F2/F3 全覆盖；7 条材料版本；4 条 Agent 运行", 4.74, 1.75, 3.85, 2.1, CYAN, "02")
    card(s, "证据边界", "项目分 68、Agent 分 71、教师判断与 D1—D9 均标为 S（模拟），不冒充真实课程成绩", 8.76, 1.75, 3.85, 2.1, AMBER, "03")
    rect(s, 0.72, 4.25, 11.89, 1.85, PANEL_2, LINE)
    text(s, "最终判定", 1.0, 4.58, 1.5, 0.35, 13, CYAN, True)
    text(s, "VentureAI 第三阶段功能与证据机制：完成", 1.0, 5.04, 6.0, 0.46, 24, TEXT, True)
    text(s, "CareAI 真实课程终局验收：仍须学生真实材料 + 任课教师复审 + 现场答辩", 7.15, 4.88, 4.95, 0.75, 15, AMBER, True)

    # 03 flow
    s = base(prs, 3, "终局验收不是一个分数，而是一条不可跳过的链", "任何关键节点失败，都不能由其他高分补偿")
    stages = [("01", "立项闯关", "G1—G6 全达标"), ("02", "材料成套", "计划书/PPT/讲稿/问答/总结"), ("03", "三流程证据", "F1/F2/F3 有真实 run_id"), ("04", "双门槛", "项目与 Agent 各 ≥60"), ("05", "终局", "完整、无红线、教师确认")]
    for i, (num, ttl, body) in enumerate(stages):
        x = 0.72 + i * 2.42
        rect(s, x, 2.05, 2.05, 2.55, PANEL, LINE)
        rect(s, x + 0.18, 2.26, 0.56, 0.56, CYAN if i < 4 else GREEN, CYAN if i < 4 else GREEN)
        text(s, num, x + 0.18, 2.26, 0.56, 0.56, 14, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(s, ttl, x + 0.18, 3.08, 1.68, 0.38, 16, TEXT, True)
        text(s, body, x + 0.18, 3.58, 1.68, 0.72, 12.5, MUTED)
        if i < 4:
            text(s, "→", x + 2.05, 3.0, 0.36, 0.5, 22, CYAN, True, PP_ALIGN.CENTER)
    rect(s, 0.72, 5.14, 11.89, 0.95, "0E2030", CYAN)
    text(s, "关键规则：立项更新、计划书更新后，旧评分和旧终局结论自动失效，必须按当前版本重新验收。", 1.0, 5.41, 11.25, 0.4, 15, TEXT, True, PP_ALIGN.CENTER)

    # 04 G gates
    s = base(prs, 4, "G1—G6：先证明项目值得做，再进入正式材料", "教师逐项判断；学生更新立项书后，旧版通过自动失效")
    gates = [
        ("G1", "问题具体", "用户、场景、问题可识别"), ("G2", "证据成立", "来源可查，结论不过度"),
        ("G3", "边界透明", "F/I/H/S 标注完整"), ("G4", "方案匹配", "回应问题并说明边界"),
        ("G5", "价值可行", "社会价值、创新、实施路径"), ("G6", "Agent 可追溯", "V2 参与且运行可定位"),
    ]
    for i, (g, ttl, body) in enumerate(gates):
        x = 0.72 + (i % 3) * 4.03; y = 1.72 + (i // 3) * 2.08
        card(s, f"{g}  {ttl}", body, x, y, 3.78, 1.72, CYAN if i < 3 else INDIGO)
    rect(s, 0.72, 6.02, 11.89, 0.62, PANEL_2, LINE)
    text(s, "再审要求：上一版 → 教师反馈 → 具体修改 → 验证结果；只换措辞不算有效修改。", 0.95, 6.19, 11.4, 0.28, 13, AMBER, True, PP_ALIGN.CENTER)

    # 05 deliverables
    s = base(prs, 5, "提交物必须成套、同源、同版本", "系统同时检查“有没有”和“是否绑定当前版本”")
    artifacts = [("01", "立项书", "通过版 + 历次反馈"), ("02", "商业计划书", "12章；关键数字可回溯"), ("03", "路演 PPT", "建议12—16页；与计划书一致"), ("04", "10分钟讲稿", "保存实际计时，≤600秒"), ("05", "5分钟问答库", "≥10个高风险问题"), ("06", "阶段总结", "成果、Agent贡献、局限")]
    for i, (num, ttl, body) in enumerate(artifacts):
        x = 0.72 + (i % 3) * 4.03; y = 1.74 + (i // 3) * 2.0
        card(s, ttl, body, x, y, 3.78, 1.62, GREEN if i in (3, 4) else CYAN, num, 13.5)
    rect(s, 0.72, 5.95, 11.89, 0.72, "0E2030", CYAN)
    text(s, "Agent 生成只能保存为草稿；必须由学生核验、修改并以 student 版本提交，才算正式材料。", 0.95, 6.15, 11.4, 0.3, 13.5, TEXT, True, PP_ALIGN.CENTER)

    # 06 F flows
    s = base(prs, 6, "F1 / F2 / F3：证明 Agent 确实参与并产生作用", "每条记录同时保存任务、输入版本、处理理由、结果版本、证据位置与效果判断")
    cards = [
        ("F1", "理论学习", "Tutor 运行\n理解检查\n错误与纠正", CYAN),
        ("F2", "项目指导", "Coach/竞品/答辩运行\n采纳·拒绝·重写·补充\n映射成果版本", INDIGO),
        ("F3", "评审反馈", "Grader 初评\n定位问题并修改\n另一 run 复验", GREEN),
    ]
    for i, (flow, ttl, body, color) in enumerate(cards):
        x = 0.72 + i * 4.03
        rect(s, x, 1.78, 3.78, 3.45, PANEL, LINE)
        text(s, flow, x + 0.28, 2.02, 0.9, 0.55, 24, color, True)
        text(s, ttl, x + 1.15, 2.07, 2.1, 0.4, 17, TEXT, True)
        rich_lines(s, body.split("\n"), x + 0.3, 2.85, 3.15, 1.65, 15, 10, True)
    rect(s, 0.72, 5.57, 11.89, 0.82, PANEL_2, LINE)
    text(s, "可追溯字段：run_id · session · project · Agent版本 · Git提交 · 模型 · 代码/提示词哈希 · 状态 · 时间", 0.95, 5.81, 11.4, 0.35, 13, MUTED, True, PP_ALIGN.CENTER)

    # 07 scoring
    s = base(prs, 7, "评分与排序：门槛负责通过，排名负责区分", "高总分不能补偿任一门槛低于60分")
    for x, value, label, color in [(0.82, "68", "项目成果 / 100\nS（模拟）", CYAN), (3.32, "71", "Agent效果 / 100\nS（模拟）", INDIGO), (5.82, "—", "现场答辩 / 100\n教师现场录入", AMBER)]:
        rect(s, x, 1.82, 2.15, 2.2, PANEL, color)
        text(s, value, x, 2.12, 2.15, 0.72, 35, color, True, PP_ALIGN.CENTER)
        text(s, label, x + 0.16, 3.02, 1.83, 0.68, 13, MUTED, True, PP_ALIGN.CENTER)
    rect(s, 8.35, 1.82, 4.15, 2.2, PANEL_2, LINE)
    text(s, "D8 暂列", 8.7, 2.12, 1.15, 0.32, 14, CYAN, True)
    text(s, "(项目分 + Agent分) ÷ 2", 8.7, 2.62, 3.1, 0.35, 19, TEXT, True)
    text(s, "仅双门槛通过者进入；约前50%", 8.7, 3.23, 3.25, 0.35, 13, MUTED)
    rect(s, 0.82, 4.45, 11.68, 1.48, "0E2030", CYAN)
    text(s, "D9 入围组最终排序", 1.16, 4.78, 2.25, 0.35, 15, CYAN, True)
    text(s, "项目分 + Agent分 + 答辩分 = 300分总成绩", 3.55, 4.68, 5.05, 0.5, 21, TEXT, True, PP_ALIGN.CENTER)
    text(s, "仍由教师确认同分、材料和红线", 9.05, 4.82, 2.95, 0.34, 13, MUTED, True, PP_ALIGN.CENTER)

    # 08 comparison
    s = base(prs, 8, "CareAI 与蚂蚁阿福：先公平比较，再谈选择", "阿福已有健康档案与健康小目标；不能靠贬低竞品建立差异化")
    cols = [0.72, 4.05, 8.35]
    widths = [3.1, 4.05, 4.25]
    headers = [("比较维度", MUTED), ("CareAI 当前证据", CYAN), ("蚂蚁阿福公开信息", GREEN)]
    for (label, color), x, w in zip(headers, cols, widths):
        rect(s, x, 1.7, w, 0.58, PANEL_2, LINE)
        text(s, label, x + 0.16, 1.87, w - 0.32, 0.25, 12, color, True)
    rows = [
        ("产品阶段", "可演示 MVP；无真实用户/留存/付费证据", "已公开运营的健康 AI 应用"),
        ("档案与目标", "基础记录、规则提示、7天行动计划", "个人/家庭健康档案、健康小目标与提醒"),
        ("服务范围", "非诊断健康教育与习惯管理", "健康咨询、报告解读、档案及挂号/云陪诊等服务"),
        ("可检验差异", "规则可追溯、AI受限、边界可见", "生态与服务能力强；具体体验须同条件实测"),
        ("证据结论", "技术功能 F；用户偏好与竞争优势 H", "功能来自官方公开资料；不推断其实际效果"),
    ]
    for i, row in enumerate(rows):
        y = 2.38 + i * 0.79
        for j, (value, x, w) in enumerate(zip(row, cols, widths)):
            rect(s, x, y, w, 0.68, PANEL if i % 2 == 0 else "0C1728", LINE, radius=False)
            text(s, value, x + 0.14, y + 0.1, w - 0.28, 0.48, 10.8 if j else 11.3, TEXT if j else MUTED, j == 0, valign=MSO_ANCHOR.MIDDLE)
    link_text(s, "来源1：蚂蚁阿福 App Store 产品页", "https://apps.apple.com/cn/app/id6743828427", 0.76, 6.43, 3.25, 0.22)
    link_text(s, "来源2：人民日报（数据来源：蚂蚁集团）", "https://paper.people.com.cn/rmrb/pc/content/202512/16/content_30124850.html", 4.15, 6.43, 4.0, 0.22)
    text(s, "核对日期 2026-09-15", 9.85, 6.43, 2.3, 0.22, 8.5, MUTED, align=PP_ALIGN.RIGHT)

    # 09 why choose
    s = base(prs, 9, "为什么用户可能选择 CareAI？答案目前是四个待验证假设", "差异化必须通过同条件任务证明，不能直接写成市场事实")
    hypotheses = [
        ("H1", "更容易理解", "规则结果与 AI 解释分层；用户能说清“这不是诊断”"),
        ("H2", "更容易行动", "固定输出“指标—原因—建议—7天计划”，减少下一步不明确"),
        ("H3", "更容易复核", "规则版本、阈值、输入和输出可追溯；错误更容易定位"),
        ("H4", "更少数据负担", "课堂场景只用虚构资料；真实发布仍须完成隐私与授权设计"),
    ]
    for i, (h, ttl, body) in enumerate(hypotheses):
        x = 0.72 + (i % 2) * 6.0; y = 1.78 + (i // 2) * 2.08
        rect(s, x, y, 5.72, 1.72, PANEL, LINE)
        rect(s, x + 0.22, y + 0.25, 0.72, 0.46, AMBER, AMBER)
        text(s, h, x + 0.22, y + 0.25, 0.72, 0.46, 13, BG, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        text(s, ttl, x + 1.15, y + 0.22, 3.8, 0.4, 17, TEXT, True)
        text(s, body, x + 0.25, y + 0.85, 5.2, 0.58, 13, MUTED)
    rect(s, 0.72, 6.1, 11.72, 0.52, "2A2112", AMBER)
    text(s, "当前可以说“值得测试”，不能说“用户已经更喜欢 CareAI”。", 0.95, 6.23, 11.25, 0.25, 14, AMBER, True, PP_ALIGN.CENTER)

    # 10 test design
    s = base(prs, 10, "公平对照测试：同一任务、同一资料、交换顺序", "不采集真实健康数据；参与人数与知情安排由课程教师决定")
    steps = [("1", "准备", "一份 S（模拟）成年人资料；两边使用相同信息"), ("2", "任务", "查看/录入资料 → 找到解释 → 找到下一步行动 → 说明边界"), ("3", "控偏", "一半先用 CareAI，一半先用阿福；记录未找到/不适用"), ("4", "判断", "结果产生后再决定哪项假设保留、修改或拒绝")]
    for i, (num, ttl, body) in enumerate(steps):
        x = 0.72 + i * 3.0
        card(s, ttl, body, x, 1.78, 2.72, 2.45, CYAN if i < 3 else GREEN, num, 12.5)
    metrics = ["任务完成", "耗时", "操作步骤", "提示理解", "行动清晰度", "边界理解", "主观偏好"]
    rect(s, 0.72, 4.65, 11.72, 1.35, PANEL_2, LINE)
    text(s, "统一观察指标", 1.0, 4.93, 1.55, 0.34, 14, CYAN, True)
    for i, m in enumerate(metrics):
        x = 2.65 + (i % 4) * 2.2; y = 4.82 + (i // 4) * 0.52
        rect(s, x, y, 1.93, 0.38, "0E2030", LINE)
        text(s, m, x, y, 1.93, 0.38, 11, TEXT, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, "输出：观察表 + 差异解释 + 局限 + 下一轮修改；不要只报平均分。", 0.86, 6.3, 11.6, 0.26, 12.5, MUTED, True, PP_ALIGN.CENTER)

    # 11-14 evidence screenshots
    screenshot_slide(prs, 11, "修改前｜项目定位被错误改写", "before-careai-coach.png", "BEFORE · P0", "把非诊断健康教育项目说成心理健康筛查，并补造 NLP、访谈和用户。", "一旦项目事实错，后续商业建议、验证任务和合规判断都会沿着错误前提扩散。", RED)
    screenshot_slide(prs, 12, "修改前｜评分有数字，却没有可靠证据边界", "before-careai-grader.png", "BEFORE · P1", "R1—R11 有分数，但部分团队、商业和材料依据不在提交文件中。", "分数无法回到材料位置，就不能证明评分正确，也不能指导学生具体修改。", RED)
    screenshot_slide(prs, 13, "修改后｜先复述事实，再说缺口", "after-careai-coach.png", "AFTER · PASS", "准确识别非诊断定位、真实完成的 MVP、规则与 AI 分层，并把未验证项说清楚。", "建议开始于学生真实材料；资料缺失时明确“未提供/待验证”，减少幻觉传播。", GREEN)
    screenshot_slide(prs, 14, "修改后｜评分绑定依据、缺口与下一步", "after-careai-grader.png", "AFTER · PASS", "R1—R11 逐项显示评分、材料依据、缺失证据和可执行建议。", "学生能定位问题，教师能复核判断，修改后还能用新 run_id 做复验闭环。", GREEN)

    # 15 current stage3 UI
    s = base(prs, 15, "修改为什么有效：把提示词约束变成产品门槛", "单靠“请勿编造”不够；数据库、权限、版本与界面共同阻断错误")
    add_picture_contain(s, STAGE3 / "student-stage3.png", 0.72, 1.68, 5.72, 3.58)
    add_picture_contain(s, STAGE3 / "teacher-stage3.png", 6.88, 1.68, 5.72, 3.58)
    text(s, "学生端：材料版本、立项状态、证据导出", 0.8, 5.38, 5.55, 0.3, 11.5, CYAN, True, PP_ALIGN.CENTER)
    text(s, "教师端：G1—G6、独立评分、D8/D9 排名", 6.96, 5.38, 5.55, 0.3, 11.5, GREEN, True, PP_ALIGN.CENTER)
    rect(s, 0.72, 5.86, 11.88, 0.7, PANEL_2, LINE)
    text(s, "事实约束 + 项目权限 + run追踪 + 材料哈希 + 版本失效 + 人工确认 = 可审计的验收系统", 0.95, 6.06, 11.4, 0.3, 13.5, TEXT, True, PP_ALIGN.CENTER)

    # 16 close
    s = base(prs, 16, "终局结论与下一步", "系统已经就绪；真实成绩仍必须来自真实材料和真实教师")
    rect(s, 0.72, 1.72, 7.55, 3.9, PANEL_2, LINE)
    text(s, "已完成", 1.05, 2.05, 1.4, 0.35, 15, GREEN, True)
    rich_lines(s, [
        "G1—G6、版本材料、F1/F2/F3、F/I/H/S、D1—D9",
        "项目 / Agent / 答辩三张100分表与 D8/D9 排名",
        "CareAI 修改前后截图、原始 JSON、权限状态码",
        "15项后端测试、TypeScript 与生产构建通过",
        "V2 标签：ventureai-v2-stage3-2026-09-15",
    ], 1.05, 2.62, 6.85, 2.48, 15, 8, True)
    rect(s, 8.62, 1.72, 3.98, 3.9, PANEL, AMBER)
    text(s, "真实验收还需", 8.98, 2.05, 2.9, 0.38, 15, AMBER, True)
    rich_lines(s, [
        "替换 S 模拟材料",
        "教师本人复审评分",
        "10分钟路演 + 5分钟答辩",
        "完成合规的对照验证",
    ], 8.98, 2.73, 3.05, 2.1, 15, 12, True)
    rect(s, 0.72, 6.02, 11.88, 0.56, GREEN, GREEN)
    text(s, "第三阶段的核心不是“AI 写了多少”，而是“结果是否可信、可追溯、可答辩”。", 0.92, 6.15, 11.5, 0.28, 15, BG, True, PP_ALIGN.CENTER)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"saved {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    # The submitted deck was shortened after review. Keep the shared styling
    # helpers and long-form builder here, but make the default command produce
    # the current compact presentation.
    from generate_stage3_ppt_compact import build_compact
    build_compact()

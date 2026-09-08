# -*- coding: utf-8 -*-
"""Render the Stage 1 Markdown submission as a self-contained Word document."""
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "第一阶段正式提交材料.md"
OUTPUT = ROOT / "第一阶段正式提交材料.docx"


def set_font(run, size=10.5, bold=False, color=None, italic=False):
    run.font.name = "微软雅黑"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._r.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), color)
    tc_pr.append(fill)


def clean_inline(text):
    # Keep source URLs visible in Word exports.  The submission specification
    # requires external evidence to remain directly locatable; stripping a
    # Markdown link down to its label makes the final Word file insufficient.
    text = re.sub(r"!?\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("**", "").replace("`", "")
    return text.replace("[", "").replace("]", "")


def add_paragraph(doc, text, style=None, indent=0, quote=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    if quote:
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.right_indent = Cm(0.5)
        shade_el = OxmlElement("w:shd")
        shade_el.set(qn("w:val"), "clear")
        shade_el.set(qn("w:color"), "auto")
        shade_el.set(qn("w:fill"), "F3F6FA")
        p._p.get_or_add_pPr().append(shade_el)
    run = p.add_run(clean_inline(text))
    set_font(run, 10.5, italic=quote)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    palette = {1: (31, 78, 121), 2: (47, 84, 150), 3: (73, 107, 158)}
    sizes = {1: 16, 2: 13, 3: 11.5}
    run = p.add_run(clean_inline(text))
    set_font(run, sizes.get(level, 10.5), bold=True, color=palette.get(level, (0, 0, 0)))
    return p


def split_row(line):
    return [clean_inline(item.strip()) for item in line.strip().strip("|").split("|")]


def add_table(doc, rows):
    data = [split_row(row) for row in rows if not re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?", row)]
    if not data:
        return
    width = max(len(row) for row in data)
    table = doc.add_table(rows=len(data), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, source_row in enumerate(data):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            value = source_row[col_index] if col_index < len(source_row) else ""
            run = cell.paragraphs[0].add_run(value)
            set_font(run, 8.5, bold=row_index == 0, color=(255, 255, 255) if row_index == 0 else None)
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            if row_index == 0:
                shade(cell, "1F4E79")
            elif row_index % 2:
                shade(cell, "F4F7FB")
    doc.add_paragraph()


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D9E2F3")
    border.append(bottom)
    p._p.get_or_add_pPr().append(border)


def build_document(lines):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.1)
        section.bottom_margin = Cm(2.1)
        section.left_margin = Cm(2.3)
        section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_table(doc, table_lines)
            continue
        if re.fullmatch(r"-{3,}", stripped):
            add_rule(doc)
        elif stripped.startswith("### "):
            add_heading(doc, stripped[4:], 3)
        elif stripped.startswith("## "):
            add_heading(doc, stripped[3:], 2)
        elif stripped.startswith("# "):
            add_heading(doc, stripped[2:], 1)
        elif stripped.startswith("> "):
            add_paragraph(doc, stripped[2:], quote=True)
        elif re.match(r"^[-*] ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_inline(stripped[2:]))
            set_font(run, 10.5)
        elif re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_inline(re.sub(r"^\d+\. ", "", stripped)))
            set_font(run, 10.5)
        else:
            add_paragraph(doc, stripped)
        index += 1
    return doc


if __name__ == "__main__":
    document = build_document(SOURCE.read_text(encoding="utf-8").splitlines())
    document.save(OUTPUT)
    print(OUTPUT)

# -*- coding: utf-8 -*-
"""Turn the user's two small evidence screenshots into two full-size slides."""
from io import BytesIO
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "第一阶段最终汇报PPT_5分钟_美化版_含证据页.pptx"
OUTPUT = ROOT / "第一阶段最终汇报PPT_5分钟_美化版_证据双页.pptx"

NAVY = RGBColor(17, 37, 64)
TEAL = RGBColor(20, 136, 137)
RED = RGBColor(191, 56, 56)
WHITE = RGBColor(255, 255, 255)
MIST = RGBColor(247, 250, 253)
SKY = RGBColor(214, 236, 247)
MINT = RGBColor(208, 237, 232)
TEXT = RGBColor(34, 45, 59)
MUTED = RGBColor(92, 105, 120)


def textbox(slide, text, x, y, w, h, size=16, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear(); frame.word_wrap = True; frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]; paragraph.alignment = align
    run = paragraph.add_run(); run.text = text
    run.font.name = "Microsoft YaHei"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    return box


def add_background(slide):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = MIST
    for x, y, size, color, transparency in [
        (11.45, 0.78, 2.05, SKY, 28), (12.35, 1.55, 1.20, MINT, 10),
        (-0.65, 6.15, 1.85, SKY, 42), (0.15, 6.55, 0.90, MINT, 18),
    ]:
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
        circle.fill.solid(); circle.fill.fore_color.rgb = color; circle.fill.transparency = transparency
        circle.line.fill.background()


def add_header(slide, number, title, subtitle):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.72))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.72), Inches(13.333), Inches(0.055))
    accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor(85, 192, 196); accent.line.fill.background()
    textbox(slide, number, 0.45, 0.08, 0.55, 0.45, 20, RGBColor(150, 212, 235), True)
    textbox(slide, title, 1.0, 0.08, 8.9, 0.45, 23, WHITE, True)
    textbox(slide, subtitle, 9.8, 0.12, 3.0, 0.35, 10, RGBColor(203, 220, 235), False, PP_ALIGN.RIGHT)
    textbox(slide, "CareAI × VentureAI  ｜  Phase 1", 9.35, 7.08, 3.3, 0.18, 8.5, MUTED, False, PP_ALIGN.RIGHT)


def add_note(slide, title, body, accent):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(6.06), Inches(12.0), Inches(0.72))
    box.fill.solid(); box.fill.fore_color.rgb = WHITE; box.line.color.rgb = RGBColor(218, 230, 240)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.65), Inches(6.06), Inches(0.09), Inches(0.72))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
    textbox(slide, title, 0.92, 6.15, 1.25, 0.18, 10.5, accent, True)
    textbox(slide, body, 2.08, 6.11, 10.2, 0.35, 12, TEXT)


def add_fitted_picture(slide, blob, ratio, x, y, w, h):
    target_ratio = w / h
    if ratio >= target_ratio:
        pic_w, pic_h = w, w / ratio
        pic_x, pic_y = x, y + (h - pic_h) / 2
    else:
        pic_h, pic_w = h, h * ratio
        pic_x, pic_y = x + (w - pic_w) / 2, y
    slide.shapes.add_picture(BytesIO(blob), Inches(pic_x), Inches(pic_y), width=Inches(pic_w), height=Inches(pic_h))


def add_evidence_slide(prs, number, title, subtitle, blob, ratio, accent, note_title, note_body):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_header(slide, number, title, subtitle)
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.68), Inches(1.02), Inches(11.98), Inches(4.76))
    frame.fill.solid(); frame.fill.fore_color.rgb = WHITE; frame.line.color.rgb = accent; frame.line.width = Pt(1.7)
    add_fitted_picture(slide, blob, ratio, 0.82, 1.16, 11.70, 4.48)
    add_note(slide, note_title, note_body, accent)


def main():
    prs = Presentation(SOURCE)
    old = prs.slides[-1]
    pictures = sorted(
        [shape for shape in old.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE],
        key=lambda shape: shape.left,
    )
    if len(pictures) < 2:
        raise RuntimeError("The final evidence slide needs two inserted screenshots before it can be split.")

    evidence = []
    for picture in pictures[:2]:
        image_size = picture.image.size
        ratio = image_size[0] / image_size[1]
        evidence.append((picture.image.blob, ratio))

    # Remove the old two-column evidence page only after reading its pictures.
    slide_id = prs.slides._sldIdLst[-1]
    prs.part.drop_rel(slide_id.rId)
    del prs.slides._sldIdLst[-1]

    add_evidence_slide(
        prs, "07", "原始运行证据：V1 首次失败", "会话 83618b05-396b-42cc-9985-34ec142752a4",
        evidence[0][0], evidence[0][1], RED,
        "你要讲的话：",
        "这是 V1 的第一次运行。它把 F/I/H/S 分类和风险审阅误判成“代写”，所以没有完成任务。",
    )
    add_evidence_slide(
        prs, "08", "原始运行证据：V1.1 修复后复测", "会话 5b369e6c-4377-4587-844c-0418014cad24",
        evidence[1][0], evidence[1][1], TEAL,
        "你要讲的话：",
        "这是独立的 V1.1 新会话。它将无来源主张标为 H、模拟反馈标为 S，并给出医疗与隐私边界。",
    )
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
